"""Bounded, literal cell reading shared by Delivery and its versioned Adapter."""

from collections.abc import Callable, Sequence
from dataclasses import dataclass, field

from openpyxl.workbook.workbook import Workbook

from deepaha.documents.parser import ExpectedParseError
from deepaha.documents.spreadsheet import (
    _ignore_declared_dimensions,
    _load_workbook,
    _preflight_archive,
    _preflight_loaded_worksheet,
)
from deepaha.documents.spreadsheet_limits import (
    MAX_MERGED_CELL_VALUES,
    MAX_WORKSHEETS,
    WorksheetExpansionBudget,
    parse_merge_range,
)


@dataclass(frozen=True, slots=True)
class SpreadsheetText:
    rows: dict[str, dict[int, dict[int, str]]]
    row_counts: dict[str, int]
    display_ambiguities: dict[str, dict[int, dict[int, tuple[str, str]]]] = field(
        default_factory=dict
    )


def read_spreadsheet_text(
    content: bytes, *, open_workbook: Callable[[bytes], Workbook] = _load_workbook
) -> SpreadsheetText:
    _preflight_archive(content)
    book = open_workbook(content)
    try:
        if len(book.sheetnames) > MAX_WORKSHEETS:
            raise ExpectedParseError("XLSX_WORKSHEET_LIMIT_EXCEEDED")
        sheets: dict[str, dict[int, dict[int, str]]] = {}
        row_counts: dict[str, int] = {}
        display_ambiguities: dict[str, dict[int, dict[int, tuple[str, str]]]] = {}
        expansion = WorksheetExpansionBudget()
        for worksheet in book.worksheets:
            merges = _preflight_loaded_worksheet(worksheet, expansion)
            _ignore_declared_dimensions(worksheet)
            rows = sheets[worksheet.title] = {}
            ambiguous_rows = display_ambiguities[worksheet.title] = {}
            row_counts[worksheet.title] = 0
            for number, cells in enumerate(worksheet.iter_rows(), 1):
                populated = {
                    i: str(cell.value) for i, cell in enumerate(cells, 1) if cell.value is not None
                }
                ambiguous = {
                    i: (cell.data_type, cell.number_format)
                    for i, cell in enumerate(cells, 1)
                    if cell.value is not None
                    and (
                        cell.data_type in {"f", "b", "e"}
                        or cell.number_format not in {"General", "@"}
                    )
                }
                if populated:
                    rows[number] = populated
                if ambiguous:
                    ambiguous_rows[number] = ambiguous
                row_counts[worksheet.title] = number
            merged_last_row = _fill_merged_cells(rows, merges)
            row_counts[worksheet.title] = max(row_counts[worksheet.title], merged_last_row)
        return SpreadsheetText(sheets, row_counts, display_ambiguities)
    finally:
        book.close()


def _fill_merged_cells(rows: dict[int, dict[int, str]], merges: Sequence[str]) -> int:
    """Give every covered cell the value Excel shows: its anchor's.

    A read-only worksheet has no ``merged_cells``, so covered cells would stay
    empty and a locator pointing inside a merged range would look like a quote
    mismatch. Existing values are never overwritten. Past the cell budget the
    merge is left unfilled, degrading to today's behavior instead of failing the
    whole workbook. Returns the highest row a merge range reaches, or 0.
    """
    highest_row = 0
    filled = 0
    for reference in merges:
        bounds = parse_merge_range(reference)
        if bounds is None:
            continue
        first_row, first_column, last_row, last_column = bounds
        highest_row = max(highest_row, last_row)
        anchor = rows.get(first_row, {}).get(first_column)
        if anchor is None or filled >= MAX_MERGED_CELL_VALUES:
            continue
        for row in range(first_row, last_row + 1):
            for column in range(first_column, last_column + 1):
                if row == first_row and column == first_column:
                    continue
                if column in rows.get(row, {}):
                    continue
                rows.setdefault(row, {})[column] = anchor
                filled += 1
                if filled >= MAX_MERGED_CELL_VALUES:
                    break
            if filled >= MAX_MERGED_CELL_VALUES:
                break
    return highest_row
