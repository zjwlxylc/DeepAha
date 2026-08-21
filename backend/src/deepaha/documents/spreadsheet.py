import json
import math
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from hashlib import sha256
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook

from deepaha.contracts.phase2 import EvidenceLocatorV02, SpreadsheetRangeLocator
from deepaha.documents.normalization import normalize_text
from deepaha.documents.parser import ExpectedParseError, ParsedDocument

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_ARCHIVE_ENTRIES = 10_000
MAX_UNCOMPRESSED_BYTES = 100_000_000


@dataclass(frozen=True, slots=True)
class NormalizedCell:
    row: int
    column: int
    value: str


class OpenpyxlSpreadsheetParser:
    name = "xlsx_openpyxl"
    version = "0.2.0"

    def supports(self, media_type: str) -> bool:
        return media_type.partition(";")[0].strip().lower() == XLSX_MEDIA_TYPE

    def parse(self, content: bytes, *, artifact_sha256: str) -> ParsedDocument:
        if sha256(content).hexdigest() != artifact_sha256:
            raise ValueError("artifact SHA-256 does not match XLSX bytes")
        _preflight_archive(content)
        workbook = _load_workbook(content)
        try:
            text_lines: list[str] = []
            locators: list[EvidenceLocatorV02] = []
            for worksheet in workbook.worksheets:
                sheet_lines: list[str] = []
                for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                    normalized_values = tuple(_normalize_cell_value(cell.value) for cell in row)
                    populated_columns = [
                        index
                        for index, value in enumerate(normalized_values, start=1)
                        if value != ""
                    ]
                    if not populated_columns:
                        continue
                    start_column = populated_columns[0]
                    end_column = populated_columns[-1]
                    cells = tuple(
                        NormalizedCell(
                            row=row_number,
                            column=column,
                            value=normalized_values[column - 1],
                        )
                        for column in range(start_column, end_column + 1)
                    )
                    locators.append(
                        SpreadsheetRangeLocator(
                            schema_version="0.2.0",
                            kind="spreadsheet_range",
                            sheet_name=worksheet.title,
                            start_row=row_number,
                            end_row=row_number,
                            start_column=start_column,
                            end_column=end_column,
                            cells_sha256=hash_normalized_cells(cells),
                        )
                    )
                    sheet_lines.extend(
                        f"{get_column_letter(column)}{row_number}\t{value}"
                        for column, value in enumerate(normalized_values, start=1)
                        if value != ""
                    )
                if sheet_lines:
                    if text_lines:
                        text_lines.append("")
                    text_lines.append(f"[{worksheet.title}]")
                    text_lines.extend(sheet_lines)
            if not locators:
                raise ExpectedParseError("XLSX_EMPTY")
            return ParsedDocument(
                title=None,
                published_at=None,
                language="und",
                normalized_text=normalize_text("\n".join(text_lines) + "\n"),
                locators=tuple(locators),
                needs_review_reasons=(),
            )
        finally:
            workbook.close()


def hash_normalized_cells(cells: Sequence[NormalizedCell]) -> str:
    payload = [{"row": cell.row, "column": cell.column, "value": cell.value} for cell in cells]
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return sha256(encoded).hexdigest()


def _preflight_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content), "r") as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ARCHIVE_ENTRIES:
                raise ExpectedParseError("XLSX_ENTRY_LIMIT_EXCEEDED")
            if sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ExpectedParseError("XLSX_UNCOMPRESSED_SIZE_EXCEEDED")

            lower_names: set[str] = set()
            for entry in entries:
                normalized_name = entry.filename.replace("\\", "/")
                path = PurePosixPath(normalized_name)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or (path.parts and ":" in path.parts[0])
                ):
                    raise ExpectedParseError("XLSX_PATH_TRAVERSAL")
                lower_names.add(normalized_name.lower())

            if any(name.endswith("vbaproject.bin") for name in lower_names):
                raise ExpectedParseError("XLSX_MACRO_NOT_ALLOWED")
            if any(name.startswith("xl/externallinks/") for name in lower_names):
                raise ExpectedParseError("XLSX_EXTERNAL_LINK_NOT_ALLOWED")

            content_types_name = next(
                (name for name in archive.namelist() if name.lower() == "[content_types].xml"),
                None,
            )
            if content_types_name is not None:
                content_types = archive.read(content_types_name).lower()
                if b"macroenabled" in content_types or b"vbaproject" in content_types:
                    raise ExpectedParseError("XLSX_MACRO_NOT_ALLOWED")
    except ExpectedParseError:
        raise
    except (BadZipFile, KeyError, OSError) as error:
        raise ExpectedParseError("XLSX_PARSE_FAILED") from error


def _load_workbook(content: bytes) -> Workbook:
    try:
        return load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, TypeError, ValueError) as error:
        raise ExpectedParseError("XLSX_PARSE_FAILED") from error


def _normalize_cell_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return normalize_text(value)
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ExpectedParseError("XLSX_CELL_VALUE_INVALID")
        return repr(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value.total_seconds())
    raise ExpectedParseError("XLSX_CELL_TYPE_UNSUPPORTED")


def _cells_for_range(
    workbook: Workbook,
    locator: SpreadsheetRangeLocator,
) -> tuple[NormalizedCell, ...]:
    if locator.sheet_name not in workbook.sheetnames:
        raise LookupError("spreadsheet locator sheet does not exist")
    worksheet = workbook[locator.sheet_name]
    cells: list[NormalizedCell] = []
    for row_offset, row in enumerate(
        worksheet.iter_rows(
            min_row=locator.start_row,
            max_row=locator.end_row,
            min_col=locator.start_column,
            max_col=locator.end_column,
        )
    ):
        cells.extend(
            NormalizedCell(
                row=locator.start_row + row_offset,
                column=locator.start_column + column_offset,
                value=_normalize_cell_value(cell.value),
            )
            for column_offset, cell in enumerate(row)
        )
    return tuple(cells)
