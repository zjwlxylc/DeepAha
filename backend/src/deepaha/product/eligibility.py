"""SG4 deterministic, evidence-backed eligibility computation lane.

Display admission and qualification are intentionally separate.  This module never
promotes WMA claims into legacy VerifiedFact/RuleSet tables and never calls a model.
INELIGIBLE is emitted only for a current, mandatory criterion whose evidence can be
located in immutable stored material and whose comparison is deterministic.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from hashlib import sha256
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from xml.etree import ElementTree as ET
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select

from .models import CatalogTarget, Opportunity, Revision, Snapshot
from .qualification_compiler import compile_field

STATUS_ELIGIBLE='ELIGIBLE'
STATUS_LIKELY='LIKELY_ELIGIBLE'
STATUS_UNCERTAIN='UNCERTAIN'
STATUS_INELIGIBLE='INELIGIBLE'

EDUCATION_LEVELS={'专科':1,'大专':1,'本科':2,'硕士研究生':3,'硕士':3,'博士研究生':4,'博士':4}
PROFILE_EDUCATION={'专科':1,'本科':2,'硕士研究生':3,'博士研究生':4}
PROFILE_DEGREE={'无':0,'学士':1,'硕士':2,'博士':3}
EXCEPTION_TOKENS=('放宽','例外','除外','另有','或具有','或取得','符合其一','满足其一')
NONEXCLUSIVE_MAJOR_TOKENS=('等','相关专业','相近专业','类似专业','不限于')
UNSUPPORTED_HARD_LABELS=('其他条件','学位','国籍','政治面貌','党员','英语','外语','工作经历','职业资格','资格证','职称','基层经历','身份要求')


def _norm(value: object) -> str:
    return ' '.join(str(value or '').replace('\r',' ').replace('\n',' ').split())


def _field_key(scope: str, label: str, occurrence: int) -> str:
    return f'{scope}:{label}#{occurrence}'


def _exception(value: str) -> bool:
    return any(token in value for token in EXCEPTION_TOKENS)


def _hard_field(label: str, value: str) -> bool:
    if any(token in label for token in ('学历','专业','年龄','户籍','生源','毕业届别','届别','招聘对象')):
        return True
    return any(token in label for token in UNSUPPORTED_HARD_LABELS)


def _candidate(field: dict, scope: str, occurrence: int, all_fields: list[dict]) -> dict | None:
    return compile_field(field,scope,occurrence,all_fields)

def _cell_text(value: Any) -> str:
    if value is None:return ''
    if isinstance(value,bool):return 'TRUE' if value else 'FALSE'
    if isinstance(value,float) and value.is_integer():return str(int(value))
    return _norm(value)


@dataclass
class EvidenceResolver:
    product: Any
    revision: Revision
    snapshot: Snapshot

    def __post_init__(self):
        content=self.revision.content if isinstance(self.revision.content,dict) else {}
        rows=content.get('artifacts') if isinstance(content.get('artifacts'),list) else []
        self.artifacts={str(x.get('id')):x for x in rows if isinstance(x,dict) and x.get('id')}
        self.by_url={}
        for aid,a in self.artifacts.items():
            if a.get('url'):self.by_url.setdefault(str(a.get('url')),[]).append((aid,a))
        self.workbooks={};self.text_cache={}

    def _raw(self, artifact: dict):
        path=artifact.get('path')
        if not isinstance(path,str) or path not in self.snapshot.manifest:return None
        try:
            raw=self.product.store.one(self.snapshot.manifest,path)
        except Exception:
            return None
        expected=self.snapshot.manifest[path].get('sha256')
        if expected and sha256(raw).hexdigest()!=expected:return None
        return raw

    def _extract_text(self, aid: str, artifact: dict, locator: dict) -> str | None:
        key=(aid,str(locator.get('page') or ''))
        if key in self.text_cache:return self.text_cache[key]
        path=str(artifact.get('path') or '')
        raw=self._raw(artifact)
        if raw is None:return None
        ext=path.lower().rsplit('.',1)[-1] if '.' in path else ''
        text=None
        try:
            if ext in {'txt','md','json','csv','xml'}:
                text=raw.decode('utf-8-sig',errors='replace')
            elif ext in {'html','htm'}:
                from .adapter import searchable
                text=searchable(raw,path)
            elif ext=='docx':
                with ZipFile(BytesIO(raw)) as z:
                    chunks=[]
                    for name in z.namelist():
                        if not (name=='word/document.xml' or name.startswith('word/header') or name.startswith('word/footer')):continue
                        root=ET.fromstring(z.read(name))
                        chunks.extend(x.text or '' for x in root.iter() if x.tag.endswith('}t'))
                    text=' '.join(chunks)
            elif ext=='pdf':
                try:
                    from pypdf import PdfReader
                    reader=PdfReader(BytesIO(raw))
                    page=locator.get('page')
                    if isinstance(page,str) and page.isdigit():page=int(page)
                    pages=[reader.pages[page-1]] if type(page) is int and 1<=page<=len(reader.pages) else reader.pages
                    text='\n'.join((pg.extract_text() or '') for pg in pages)
                except Exception:
                    text=None
        except Exception:
            text=None
        text=_norm(text) if text is not None else None
        self.text_cache[key]=text
        return text

    def _derived_text_match(self, artifact_id: str, artifact: dict, quote: str, locator: dict):
        # A derived text/OCR artifact is accepted only when both original and
        # derived bytes are present and integrity-matched, and they share the exact
        # source URL.  This preserves lineage without running OCR at eligibility time.
        if self._raw(artifact) is None or not artifact.get('url'):return None
        for did,d in self.by_url.get(str(artifact.get('url')),[]):
            if did==artifact_id or d.get('integrity')!='MATCH':continue
            path=str(d.get('path') or '').lower()
            if not path.endswith(('.txt','.md','.json','.html','.htm','.xml','.csv')):continue
            text=self._extract_text(did,d,{})
            if text and _norm(quote) in text:return did
        return None

    def verify(self, ref: dict) -> dict:
        result={'artifact_id':str(ref.get('artifact_id') or ''),'quote':_norm(ref.get('quote')),
                'locator':ref.get('locator') if isinstance(ref.get('locator'),dict) else {},
                'url':ref.get('url'),'verification':'UNLOCATED','located':False}
        artifact=self.artifacts.get(result['artifact_id'])
        if not artifact or artifact.get('integrity')!='MATCH' or not result['quote']:
            return result
        path=str(artifact.get('path') or '')
        locator=result['locator'];sheet=locator.get('sheet');row=locator.get('row')
        column=locator.get('column') or locator.get('col');cell=locator.get('cell')
        if isinstance(row,str) and row.isdigit():row=int(row)
        # XLSX cell/range verification remains exact.
        if path.lower().endswith('.xlsx') and path in self.snapshot.manifest and isinstance(sheet,str) and type(row) is int and row>=1:
            try:
                if isinstance(cell,str):
                    m=re.fullmatch(r'([A-Za-z]{1,3})([1-9][0-9]{0,6})',cell.strip())
                    if not m:return result
                    column=m.group(1).upper();row=int(m.group(2))
                if isinstance(column,str):column=column_index_from_string(column.upper())
                if column is not None and (type(column) is not int or column<1):return result
                book=self.workbooks.get(path)
                if book is None:
                    raw=self._raw(artifact)
                    if raw is None:return result
                    book=load_workbook(BytesIO(raw),read_only=True,data_only=True,keep_links=False);self.workbooks[path]=book
                if sheet not in book.sheetnames:return result
                ws=book[sheet];quote=result['quote']
                if column is not None:matched=_cell_text(ws.cell(row=row,column=column).value)==quote
                else:matched=quote in {_cell_text(c.value) for c in next(ws.iter_rows(min_row=row,max_row=row))}
                if matched:result.update(verification='XLSX_CELL_MATCH',located=True)
                return result
            except (OSError,ValueError,KeyError,IndexError,TypeError,BadZipFile,InvalidFileException):return result
        # Adapter-located text was already matched against immutable stored bytes at intake.
        if ref.get('located') is True:
            result.update(verification='TEXT_QUOTE_MATCH',located=True);return result
        text=self._extract_text(result['artifact_id'],artifact,locator)
        if text and result['quote'] in text:
            ext=path.lower().rsplit('.',1)[-1] if '.' in path else ''
            result.update(verification={'docx':'DOCX_TEXT_MATCH','pdf':'PDF_TEXT_MATCH'}.get(ext,'TEXT_QUOTE_MATCH'),located=True);return result
        derived=self._derived_text_match(result['artifact_id'],artifact,result['quote'],locator)
        if derived:
            result.update(verification='DERIVED_TEXT_MATCH',located=True,derived_artifact_id=derived)
        return result


def _support(candidate: dict, resolver: EvidenceResolver) -> tuple[str,list[dict],str | None]:
    state=candidate['state']
    if state in {'CONFLICT','INVALID','PROVEN_WRONG','RETRACTED','EXCLUDED'}:
        return 'CONFLICTING_OR_INVALID',[], 'SOURCE_STATE_NOT_COMPUTABLE'
    if state not in {'CONFIRMED','KNOWN'}:
        return 'UNLOCATED',[], 'SOURCE_STATE_NOT_CONFIRMED'
    verified=[resolver.verify(x) for x in candidate['evidence'] if isinstance(x,dict)]
    located=[x for x in verified if x['located']]
    if candidate.get('support_mode')=='ALL':
        required=int(candidate.get('required_evidence_count') or len(verified) or 1)
        if len(verified)<required or len(located)!=len(verified):
            return 'UNLOCATED',verified,'EVIDENCE_UNLOCATED'
    elif not located:
        return 'UNLOCATED',verified,'EVIDENCE_UNLOCATED'
    return 'LOCATED',verified,None


def _evaluate(candidate: dict, profile: dict) -> tuple[str,str,dict]:
    typ=candidate['type'];n=candidate['normalized']
    if not candidate['deterministic']:
        return 'UNKNOWN',candidate['normalization_reason'],{}
    if n.get('no_restriction'):
        return 'SATISFIED','NO_RESTRICTION',{}
    if typ=='EDUCATION_MIN':
        value=profile.get('education');rank=PROFILE_EDUCATION.get(value)
        if rank is None:return 'UNKNOWN','PROFILE_EDUCATION_MISSING',{'profile_value':value}
        return ('SATISFIED','MEETS_MINIMUM',{'profile_value':value}) if rank>=n['minimum_level'] else ('CONFLICT','BELOW_MINIMUM_EDUCATION',{'profile_value':value})
    if typ=='GRADUATION_YEAR_SET':
        year=profile.get('graduation_year')
        if not isinstance(year,int):return 'UNKNOWN','PROFILE_GRADUATION_YEAR_MISSING',{'profile_value':year}
        return ('SATISFIED','GRADUATION_YEAR_ALLOWED',{'profile_value':year}) if year in n['allowed_years'] else ('CONFLICT','GRADUATION_YEAR_CONFLICT',{'profile_value':year})
    if typ=='MAJOR_CODE_SET':
        code=str(profile.get('major_code') or '')
        if not re.fullmatch(r'\d{2,8}',code):return 'UNKNOWN','PROFILE_MAJOR_CODE_MISSING',{'profile_value':code}
        if any(code.startswith(prefix) for prefix in n['allowed_prefixes']):return 'SATISFIED','MAJOR_CODE_ALLOWED',{'profile_value':code}
        if n.get('exclusive',True):return 'CONFLICT','MAJOR_CODE_CONFLICT',{'profile_value':code}
        return 'UNKNOWN','MAJOR_LIST_NONEXCLUSIVE',{'profile_value':code}
    if typ=='BIRTH_DATE_MIN':
        raw=profile.get('birth_date')
        try:birth=date.fromisoformat(raw) if isinstance(raw,str) else None
        except ValueError:birth=None
        if birth is None:return 'UNKNOWN','PROFILE_BIRTH_DATE_MISSING',{'profile_value':raw}
        cutoff=date.fromisoformat(n['minimum_birth_date'])
        return ('SATISFIED','BIRTH_DATE_WITHIN_LIMIT',{'profile_value':raw}) if birth>=cutoff else ('CONFLICT','AGE_LIMIT_CONFLICT',{'profile_value':raw})
    if typ=='HUKOU_REGION':
        region=_norm(profile.get('hukou_region'))
        if not region:return 'UNKNOWN','PROFILE_HUKOU_REGION_MISSING',{'profile_value':region}
        if any(req in region for req in n['allowed_regions']):return 'SATISFIED','HUKOU_REGION_ALLOWED',{'profile_value':region}
        return 'CONFLICT','HUKOU_REGION_CONFLICT',{'profile_value':region}
    if typ=='DEGREE_MIN':
        value=profile.get('degree');rank=PROFILE_DEGREE.get(value)
        if rank is None:return 'UNKNOWN','PROFILE_DEGREE_MISSING',{'profile_value':value}
        return ('SATISFIED','MEETS_MINIMUM_DEGREE',{'profile_value':value}) if rank>=n['minimum_level'] else ('CONFLICT','BELOW_MINIMUM_DEGREE',{'profile_value':value})
    if typ=='NATIONALITY':
        value=_norm(profile.get('nationality'))
        if not value:return 'UNKNOWN','PROFILE_NATIONALITY_MISSING',{'profile_value':value}
        allowed=n.get('allowed') or []
        ok=any(x in value or value in x for x in allowed)
        return ('SATISFIED','NATIONALITY_ALLOWED',{'profile_value':value}) if ok else ('CONFLICT','NATIONALITY_CONFLICT',{'profile_value':value})
    if typ=='POLITICAL_STATUS':
        value=_norm(profile.get('political_status'))
        if not value:return 'UNKNOWN','PROFILE_POLITICAL_STATUS_MISSING',{'profile_value':value}
        return ('SATISFIED','POLITICAL_STATUS_ALLOWED',{'profile_value':value}) if value in n.get('allowed',[]) else ('CONFLICT','POLITICAL_STATUS_CONFLICT',{'profile_value':value})
    if typ=='EXPERIENCE_MIN_YEARS':
        value=profile.get('work_experience_years')
        if not isinstance(value,(int,float)):return 'UNKNOWN','PROFILE_EXPERIENCE_MISSING',{'profile_value':value}
        return ('SATISFIED','EXPERIENCE_MEETS_MINIMUM',{'profile_value':value}) if float(value)>=float(n['minimum_years']) else ('CONFLICT','EXPERIENCE_BELOW_MINIMUM',{'profile_value':value})
    if typ=='FRESH_GRADUATE_REQUIRED':
        value=_norm(profile.get('student_status'))
        if not value:return 'UNKNOWN','PROFILE_STUDENT_STATUS_MISSING',{'profile_value':value}
        if value in {'应届毕业生','应届生'}:return 'SATISFIED','FRESH_GRADUATE_CONFIRMED',{'profile_value':value}
        if value in {'非应届','往届','已就业'}:return 'CONFLICT','FRESH_GRADUATE_CONFLICT',{'profile_value':value}
        return 'UNKNOWN','PROFILE_STUDENT_STATUS_AMBIGUOUS',{'profile_value':value}
    if typ=='BIRTH_DATE_BY_EDUCATION':
        raw=profile.get('birth_date');rank=PROFILE_EDUCATION.get(profile.get('education'))
        try:birth=date.fromisoformat(raw) if isinstance(raw,str) else None
        except ValueError:birth=None
        cutoff=(n.get('minimum_birth_dates') or {}).get(str(rank)) if rank else None
        if birth is None:return 'UNKNOWN','PROFILE_BIRTH_DATE_MISSING',{'profile_value':raw}
        if not cutoff:return 'UNKNOWN','AGE_RULE_NOT_AVAILABLE_FOR_EDUCATION',{'profile_value':raw}
        return ('SATISFIED','BIRTH_DATE_WITHIN_LIMIT',{'profile_value':raw}) if birth>=date.fromisoformat(cutoff) else ('CONFLICT','AGE_LIMIT_CONFLICT',{'profile_value':raw})
    if typ=='PARTICIPANT_SCOPE':
        rank=PROFILE_EDUCATION.get(profile.get('education'))
        if rank is None:return 'UNKNOWN','PROFILE_EDUCATION_MISSING',{'profile_value':profile.get('education')}
        levels=n.get('education_levels') or []
        return ('SATISFIED','PARTICIPANT_EDUCATION_ALLOWED',{'profile_value':profile.get('education')}) if rank in levels else ('CONFLICT','PARTICIPANT_EDUCATION_CONFLICT',{'profile_value':profile.get('education')})
    if typ=='TEAM_SIZE_MAX':
        value=profile.get('team_size')
        if not isinstance(value,int):return 'UNKNOWN','PROFILE_TEAM_SIZE_MISSING',{'profile_value':value}
        return ('SATISFIED','TEAM_SIZE_ALLOWED',{'profile_value':value}) if value<=n['maximum'] else ('CONFLICT','TEAM_SIZE_CONFLICT',{'profile_value':value})
    if typ in {'CERTIFICATE_REQUIRED','TITLE_REQUIRED','LANGUAGE_CERTIFICATE'}:
        key={'CERTIFICATE_REQUIRED':'certificates','TITLE_REQUIRED':'professional_title','LANGUAGE_CERTIFICATE':'language_certificates'}[typ]
        value=profile.get(key)
        if value in ('',None,[]):return 'UNKNOWN','PROFILE_REQUIREMENT_MISSING',{'profile_value':value}
        values=value if isinstance(value,list) else [value];required=_norm(n.get('required_text'))
        if any(_norm(v) and (_norm(v) in required or required in _norm(v)) for v in values):return 'SATISFIED','PROFILE_REQUIREMENT_MATCH',{'profile_value':value}
        # A profile list is not assumed exhaustive; mismatch never hard-denies.
        return 'UNKNOWN','PROFILE_REQUIREMENT_NOT_RECORDED',{'profile_value':value}
    return 'UNKNOWN','UNSUPPORTED_HARD_CONDITION',{}



def preflight_qualification_risk(content: dict, profile: dict) -> dict:
    """Return a conservative recommendation-only qualification risk signal.

    This helper deliberately ignores evidence *for the purpose of ranking only*.
    It never changes formal SG4 eligibility.  A HIGH result means: a candidate
    condition can be deterministically normalized and conflicts with the
    supplied profile, but callers must still wait for the evidence-backed
    evaluator before emitting INELIGIBLE.
    """
    flattened=[]
    all_raw=[]
    for scope in ('fields','ancestor_fields','common_fields'):
        counts={}
        for field in content.get(scope,[]) if isinstance(content.get(scope),list) else []:
            if not isinstance(field,dict):
                continue
            label=_norm(field.get('label'))
            counts[label]=counts.get(label,0)+1
            flattened.append((scope,field,counts[label]))
            all_raw.append(field)
    conflicts=[]
    for scope,field,occ in flattened:
        cand=_candidate(field,scope,occ,all_raw)
        if not cand or not cand.get('deterministic') or cand.get('type')=='APPLICATION_DEADLINE':
            continue
        outcome,reason,_profile_view=_evaluate(cand,profile)
        if outcome=='CONFLICT':
            conflicts.append({
                'type':cand.get('type'),'label':cand.get('label'),
                'reason_code':reason,'field_key':cand.get('field_key'),
            })
    return {'level':'HIGH' if conflicts else 'NORMAL','potential_conflicts':conflicts}

def _deadline_candidate(content: dict, fields: list[tuple[str,dict,int]]) -> dict | None:
    day=content.get('deadline')
    if not isinstance(day,str):return None
    try:date.fromisoformat(day)
    except ValueError:return None
    evidence=[];source_field=None
    iso=day;cn=f'{int(day[:4])}年{int(day[5:7])}月{int(day[8:10])}日'
    for scope,field,occ in fields:
        label=_norm(field.get('label'));value=_norm(field.get('value'))
        if not any(token in label for token in ('报名截止','截止时间','报名时间')):continue
        if iso not in value and cn not in value:continue
        evidence=field.get('evidence') if isinstance(field.get('evidence'),list) else []
        source_field=(scope,field,occ);break
    if source_field is None:return None
    scope,field,occ=source_field
    return {
        'field_id':field.get('id'),'field_key':_field_key(scope,_norm(field.get('label')),occ),
        'label':'报名窗口','raw_value':_norm(field.get('value')),'scope':scope,'hard':True,
        'state':str(field.get('state') or 'UNKNOWN').upper(),'evidence':evidence,
        'normalized':{'deadline':day},'type':'APPLICATION_DEADLINE','deterministic':True,'normalization_reason':'NORMALIZED',
    }


def evaluate(product: Any, target: CatalogTarget, opportunity: Opportunity, profile: dict) -> dict:
    content=product._target(target,opportunity)
    revision=product.db_row(Revision,target.revision_id) if hasattr(product,'db_row') else None
    # Use a short read transaction for immutable revision/snapshot metadata.
    with product.db.tx(False) as s:
        revision=s.get(Revision,target.revision_id)
        snapshot=s.get(Snapshot,revision.snapshot_id) if revision else None
    if revision is None or snapshot is None:
        return _result(content,[],['EVIDENCE_CONTEXT_MISSING'],STATUS_UNCERTAIN,profile,'BOUNDED',[])
    resolver=EvidenceResolver(product,revision,snapshot)
    flattened=[]
    all_raw=[]
    for scope in ('fields','ancestor_fields','common_fields'):
        counts={}
        for field in content.get(scope,[]) if isinstance(content.get(scope),list) else []:
            if not isinstance(field,dict):continue
            label=_norm(field.get('label'));counts[label]=counts.get(label,0)+1
            flattened.append((scope,field,counts[label]));all_raw.append(field)
    candidates=[]
    for scope,field,occ in flattened:
        cand=_candidate(field,scope,occ,all_raw)
        if cand:candidates.append(cand)
    deadline=_deadline_candidate(content,flattened)
    if deadline:candidates.append(deadline)

    affected=set(((content.get('currentness') or {}).get('affected_fields') or []) if isinstance(content.get('currentness'),dict) else [])
    full_pending=target.status=='UPDATE_PENDING' and (content.get('currentness') or {}).get('change_kind') in {'MISSING_PENDING','WITHDRAWAL_PENDING'}
    requirements=[];reasons=[];hard_conflicts=[];potential_conflicts=[]
    for cand in candidates:
        support,evidence,support_reason=_support(cand,resolver)
        outcome='UNKNOWN';reason=support_reason or cand['normalization_reason'];profile_view={}
        field_pending=full_pending or cand['field_key'] in affected
        if target.status=='UPDATE_PENDING' and field_pending:
            reason='TARGET_UPDATE_PENDING';reasons.append('TARGET_UPDATE_PENDING')
        elif support!='LOCATED':
            if support_reason:reasons.append(support_reason)
        elif not cand['deterministic']:
            reasons.append(cand['normalization_reason'])
        elif cand['type']=='APPLICATION_DEADLINE':
            deadline_day=date.fromisoformat(cand['normalized']['deadline'])
            if deadline_day < date.today():outcome='CONFLICT';reason='APPLICATION_WINDOW_CLOSED'
            else:outcome='SATISFIED';reason='APPLICATION_WINDOW_OPEN'
        else:
            outcome,reason,profile_view=_evaluate(cand,profile)
            if outcome=='UNKNOWN':
                reasons.append(reason)
                if cand['type']=='FRESH_GRADUATE_REQUIRED' and reason=='PROFILE_STUDENT_STATUS_MISSING':
                    # Preserve SG4's prior diagnostic contract: a bare “应届毕业生”
                    # cannot be inferred from graduation_year alone.
                    reasons.append('SEMANTIC_NOT_DETERMINISTIC')
        if support!='LOCATED' and cand['deterministic'] and not field_pending:
            po,pr,pv=_evaluate(cand,profile)
            if po=='CONFLICT':potential_conflicts.append({'type':cand['type'],'label':cand['label'],'reason_code':pr,'field_key':cand['field_key']})
        row={k:cand[k] for k in ('field_id','field_key','label','raw_value','scope','type','hard','normalized','deterministic','normalization_reason')}
        row.update(outcome=outcome,reason_code=reason,evidence_support=support,evidence=evidence,profile=profile_view)
        requirements.append(row)
        if outcome=='CONFLICT' and support=='LOCATED' and cand['deterministic'] and not field_pending:
            hard_conflicts.append({'type':cand['type'],'label':cand['label'],'reason_code':reason,'evidence':evidence,'field_key':cand['field_key']})

    if target.status=='UPDATE_PENDING':reasons.append('TARGET_UPDATE_PENDING')
    coverage_raw=str(content.get('qualification_coverage') or '').upper()
    source_requirements=[x for x in requirements if x['type']!='APPLICATION_DEADLINE']
    compiler_unresolved=[x for x in source_requirements if x['evidence_support']!='LOCATED' or not x['deterministic']]
    compiler_state='NONE' if not source_requirements else ('COMPLETE' if coverage_raw=='COMPLETE' and not compiler_unresolved else 'PARTIAL')
    coverage='COMPLETE' if compiler_state=='COMPLETE' else 'BOUNDED'
    unknown=any(x['outcome']=='UNKNOWN' for x in requirements)
    if full_pending:
        status=STATUS_UNCERTAIN;hard_conflicts=[]
    elif hard_conflicts:
        status=STATUS_INELIGIBLE
    elif not requirements or unknown:
        status=STATUS_UNCERTAIN
    elif compiler_state=='COMPLETE':
        status=STATUS_ELIGIBLE
    else:
        status=STATUS_LIKELY
    return _result(content,requirements,reasons,status,profile,coverage,hard_conflicts,compiler_state,compiler_unresolved,potential_conflicts)


def _result(content,requirements,reasons,status,profile,coverage,hard_conflicts,compiler_state='NONE',compiler_unresolved=None,potential_conflicts=None):
    titles={
        STATUS_ELIGIBLE:'当前条件明确符合',STATUS_LIKELY:'当前已知条件基本符合',
        STATUS_UNCERTAIN:'暂不能完整判断',STATUS_INELIGIBLE:'发现明确不符合条件',
    }
    messages={
        STATUS_ELIGIBLE:'在本次明确覆盖的硬条件内，你填写的信息均符合。提交前仍应查看最新官方材料。',
        STATUS_LIKELY:'当前可计算的硬条件未发现冲突，但材料没有声明资格条件覆盖完整。',
        STATUS_UNCERTAIN:'还有条件无法安全自动判断，不代表你不符合；请结合未确定项和官方原文确认。',
        STATUS_INELIGIBLE:'至少一项当前官方材料支持的明确硬条件与你填写的信息冲突。',
    }
    return {
        'status':status,'title':titles[status],'message':messages[status],
        'decision_basis':'DETERMINISTIC_EVIDENCE_BACKED','llm_used':False,
        'coverage':{'state':coverage,'explicit':coverage=='COMPLETE','compiler_state':compiler_state,'unresolved_count':len(compiler_unresolved or [])},
        'requirements':requirements,'typed_facts':requirements,
        'hard_conflicts':hard_conflicts,'review_reasons':list(dict.fromkeys(reasons)),
        'unresolved_conditions':[{'field_key':x.get('field_key'),'label':x.get('label'),'raw_value':x.get('raw_value'),'type':x.get('type'),'reason_code':x.get('reason_code'),'evidence_support':x.get('evidence_support')} for x in (compiler_unresolved or [])],
        'qualification_risk':{'level':'HIGH' if potential_conflicts else 'NORMAL','potential_conflicts':potential_conflicts or []},
        'profile_basis':{k:profile.get(k) for k in ('education','degree','major','major_code','graduation_year','student_status','birth_date','hukou_region','nationality','political_status','work_experience_years','certificates','professional_title','language_certificates','team_size') if profile.get(k) not in ('',None,[])},
        'children':[],'official_url':content.get('official_url'),'target_id':content.get('id'),
    }
