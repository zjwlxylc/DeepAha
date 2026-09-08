"""Validate recovered WMA bytes without network, database, or model-authored file reads.

The two schemas are byte-for-byte copies from deepaha-agent-poc's
skills/deepaha-opportunity-investigator/schemas. Mechanical checks establish byte,
locator and entity consistency, never semantic truth or human approval.
"""

import json
import math
import re
import unicodedata
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any, NoReturn
from urllib.parse import unquote, urlsplit

from jsonschema import Draft7Validator
from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from pypdf import PdfReader

from deepaha.documents.html import _parse_html_tree
from deepaha.documents.parser import ExpectedParseError
from deepaha.documents.pdf import MAX_PDF_PAGES
from deepaha.documents.spreadsheet import (
    _ignore_declared_dimensions,
    _preflight_archive,
    _preflight_loaded_worksheet,
)
from deepaha.documents.spreadsheet_limits import MAX_WORKSHEETS, WorksheetExpansionBudget

SCHEMA_SHA256 = {
    "opportunities": "fbbf3f83bd078ecffaa52ab8aa79743ada115c5d88918c3a5ea25245e0f6ccd5",
    "evidence": "ce6ecc0afa3edcc756460741f082b8b0438808e2159d7cdaa15fb6c312d7266a",
}
_FILES = {"opportunities.json", "evidence.json", "report.md"}
_HASH = re.compile(r"[a-fA-F0-9]{64}\Z")
_ARTIFACT_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9._-]{0,255}\Z")
_RESERVED = re.compile(r"(?:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?\Z", re.I)
HTML_QUOTE_CANONICALIZATION_VERSION = "html-quote-c14n/1"
_HAN = r"[\u3400-\u4dbf\u4e00-\u9fff]"
_HAN_FORMAT_MARKS = re.compile(rf"(?<={_HAN})[\u200b\ufeff]+(?={_HAN})")
_HTML_TEXT_BOUNDARIES = frozenset(
    [
        "address",
        "article",
        "aside",
        "blockquote",
        "body",
        "br",
        "caption",
        "dd",
        "details",
        "dialog",
        "div",
        "dl",
        "dt",
        "fieldset",
        "figcaption",
        "figure",
        "footer",
        "form",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "hgroup",
        "hr",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "summary",
        "table",
        "tbody",
        "td",
        "tfoot",
        "th",
        "thead",
        "tr",
        "ul",
    ]
)


class DeliveryValidationError(ValueError):
    def __init__(self, code: str) -> None:
        self.code = code
        super().__init__(code)


@dataclass(frozen=True, slots=True)
class _SpreadsheetText:
    rows: dict[str, dict[int, dict[int, str]]]
    row_counts: dict[str, int]


@dataclass(frozen=True, slots=True)
class ValidatedArtifact:
    artifact_id: str
    url: str
    remote_path: str
    sha256: str
    media_type: str
    content: bytes


@dataclass(frozen=True, slots=True)
class DeliveryEvidence:
    artifact_id: str
    quote: str
    locator: dict[str, Any]
    sha256: str
    mechanically_verified: bool


@dataclass(frozen=True, slots=True)
class DeliveryFact:
    entity_id: str
    field: str
    value: str | None
    status: str
    evidence: tuple[DeliveryEvidence, ...]
    note: str | None = None


@dataclass(frozen=True, slots=True)
class ValidatedDelivery:
    opportunities: dict[str, Any]
    evidence: dict[str, Any]
    report: str
    artifacts: tuple[ValidatedArtifact, ...]
    facts: tuple[DeliveryFact, ...]
    issues: tuple[str, ...]
    sha256: str


def _fail(code: str) -> NoReturn:
    raise DeliveryValidationError(code)


def _object_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            _fail("DELIVERY_JSON_INVALID")
        result[key] = value
    return result


def _json(content: bytes) -> dict[str, Any]:
    try:
        value = json.loads(
            content.decode("utf-8"),
            object_pairs_hook=_object_pairs,
            parse_constant=lambda _: _fail("DELIVERY_JSON_INVALID"),
            parse_float=_finite_float,
        )
        if not isinstance(value, dict):
            _fail("DELIVERY_JSON_INVALID")
        return dict(value)
    except UnicodeError, ValueError, RecursionError:
        raise DeliveryValidationError("DELIVERY_JSON_INVALID") from None


def _finite_float(value: str) -> float:
    result = float(value)
    if not math.isfinite(result):
        _fail("DELIVERY_JSON_INVALID")
    return result


def _load_schema(name: str) -> dict[str, Any]:
    raw = (Path(__file__).parent / "schemas" / f"{name}.schema.json").read_bytes()
    if sha256(raw).hexdigest() != SCHEMA_SHA256[name]:
        _fail("DELIVERY_SCHEMA_INTEGRITY_ERROR")
    return _json(raw)


def _schema(name: str, value: dict[str, Any]) -> None:
    if next(Draft7Validator(_load_schema(name)).iter_errors(value), None) is not None:
        _fail("DELIVERY_SCHEMA_INVALID")


def _string(value: Any, code: str) -> str:
    if not isinstance(value, str) or not value.strip() or "\x00" in value:
        _fail(code)
    return str(value)


def _safe_path(value: Any, *, remote: bool = False) -> str:
    path = _string(value, "ARTIFACT_PATH_INVALID")
    decoded = unquote(path)
    parts = decoded.split("/")
    if (
        "\\" in decoded
        or ":" in decoded
        or any(character in path for character in '%<>"|?*')
        or any(ord(c) < 32 for c in decoded)
        or any(
            part in {".", ".."} or part.endswith((".", " ")) or _RESERVED.fullmatch(part)
            for part in parts
            if part
        )
        or "//" in decoded
        or (not remote and decoded.startswith("/"))
    ):
        _fail("ARTIFACT_PATH_INVALID")
    if not remote and (len(parts) < 2 or parts[0] != "artifacts"):
        _fail("ARTIFACT_PATH_INVALID")
    if remote and "artifacts" not in parts[:-1]:
        _fail("ARTIFACT_PATH_INVALID")
    if not parts[-1]:
        _fail("ARTIFACT_PATH_INVALID")
    return path


def _url(value: Any) -> str:
    url = _string(value, "ARTIFACT_URL_INVALID")
    try:
        parsed = urlsplit(url)
        if (
            parsed.scheme not in {"http", "https"}
            or not parsed.hostname
            or parsed.username is not None
            or parsed.password is not None
        ):
            _fail("ARTIFACT_URL_INVALID")
    except ValueError:
        raise DeliveryValidationError("ARTIFACT_URL_INVALID") from None
    return url


def _manifest(evidence: dict[str, Any]) -> list[tuple[dict[str, Any], str, str, str]]:
    result: list[tuple[dict[str, Any], str, str, str]] = []
    ids: set[str] = set()
    names: set[str] = set()
    paths: set[str] = set()
    for item in evidence["artifacts"]:
        identity = _string(item.get("artifact_id"), "ARTIFACT_ID_INVALID")
        if not _ARTIFACT_ID.fullmatch(identity):
            _fail("ARTIFACT_ID_INVALID")
        if identity in ids:
            _fail("ARTIFACT_ID_DUPLICATE")
        ids.add(identity)
        local = _safe_path(item.get("local_path"))
        remote = _safe_path(item.get("remote_path", local), remote=True)
        name = PurePosixPath(unquote(remote)).name
        declared_name = item.get("file_name", name)
        if declared_name != name or PurePosixPath(local).name != name:
            _fail("ARTIFACT_PATH_INVALID")
        name_key = unicodedata.normalize("NFC", name).casefold()
        path_key = unicodedata.normalize("NFC", unquote(remote)).casefold()
        if name_key in names or path_key in paths:
            _fail("ARTIFACT_PATH_COLLISION")
        names.add(name_key)
        paths.add(path_key)
        declared = item.get("sha256")
        if not isinstance(declared, str) or not _HASH.fullmatch(declared):
            _fail("ARTIFACT_HASH_INVALID")
        _url(item.get("source_url"))
        _string(item.get("media_type"), "ARTIFACT_METADATA_INVALID")
        result.append((item, identity, remote, local))
    if not result:
        _fail("ARTIFACT_MANIFEST_EMPTY")
    return result


def _artifacts(
    evidence: dict[str, Any], contents: dict[str, bytes]
) -> tuple[ValidatedArtifact, ...]:
    result: list[ValidatedArtifact] = []
    for item, identity, remote, _local in _manifest(evidence):
        content = contents.get(identity)
        if not isinstance(content, bytes) or not content:
            _fail("ARTIFACT_MISSING")
        digest = sha256(content).hexdigest()
        if digest != item["sha256"].lower():
            _fail("ARTIFACT_HASH_MISMATCH")
        result.append(
            ValidatedArtifact(
                identity,
                _url(item.get("source_url")),
                remote,
                digest,
                _string(item.get("media_type"), "ARTIFACT_METADATA_INVALID"),
                content,
            )
        )
    if set(contents) != {item.artifact_id for item in result}:
        _fail("ARTIFACT_SET_MISMATCH")
    return tuple(result)


def _entities(o: dict[str, Any], e: dict[str, Any]) -> tuple[dict[str, Any], str]:
    indexed: dict[str, Any] = {}
    for entity in e["entities"]:
        identity = _string(entity["id"], "ENTITY_ID_INVALID")
        if identity in indexed:
            _fail("ENTITY_ID_DUPLICATE")
        indexed[identity] = entity
    roots = [value for value in indexed.values() if value["kind"] == "announcement"]
    if len(roots) != 1 or roots[0].get("parent_id") is not None:
        _fail("ENTITY_ROOT_INVALID")
    root = roots[0]
    if root["name"] != o["opportunity_name"]:
        _fail("ENTITY_REPRESENTATION_MISMATCH")
    represented = {root["id"]}
    for unit in o["units"]:
        _entity_match(unit, indexed, "unit", root["id"], represented)
        for position in unit["positions"]:
            _entity_match(position, indexed, "position", unit["id"], represented)
    if represented != set(indexed):
        _fail("ENTITY_REPRESENTATION_MISMATCH")
    return indexed, str(root["id"])


def _entity_match(
    node: dict[str, Any], indexed: dict[str, Any], kind: str, parent: str, represented: set[str]
) -> None:
    identity = node["id"]
    if identity in represented:
        _fail("ENTITY_ID_DUPLICATE")
    represented.add(identity)
    entity = indexed.get(identity)
    if entity is None or entity["kind"] != kind or entity["name"] != node["name"]:
        _fail("ENTITY_REPRESENTATION_MISMATCH")
    # The POC schema uses null for a parent already expressed by nested containment.
    nested_parent = node.get("parent_id")
    if entity.get("parent_id") != parent or nested_parent not in {None, parent}:
        _fail("ENTITY_PARENT_MISMATCH")
    if "code" in node and entity.get("code") != node["code"]:
        _fail("ENTITY_REPRESENTATION_MISMATCH")


def _fact_key(entity: str, fact: dict[str, Any]) -> tuple[str, str]:
    return entity, _string(fact.get("field"), "FACT_FIELD_INVALID")


def _fact_payload(fact: dict[str, Any]) -> str:
    value = {key: fact.get(key) for key in ("field", "value", "status", "note")}
    value["evidence"] = sorted(
        (
            {
                "artifact_id": ref.get("artifact_id"),
                "quote": ref.get("quote"),
                "locator": ref.get("locator", {}),
                "sha256": ref.get("sha256"),
            }
            for ref in fact.get("evidence", [])
        ),
        key=lambda ref: json.dumps(ref, sort_keys=True),
    )
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def _facts(
    o: dict[str, Any], e: dict[str, Any], entities: dict[str, Any], root: str
) -> list[tuple[str, dict[str, Any]]]:
    nested: dict[tuple[str, str], dict[str, Any]] = {}

    def add(entity: str, facts: list[dict[str, Any]]) -> None:
        for fact in facts:
            key = _fact_key(entity, fact)
            if key in nested:
                _fail("FACT_FIELD_DUPLICATE")
            nested[key] = fact

    add(root, o["announcement_level"])
    for unit in o["units"]:
        add(unit["id"], unit.get("unit_level", []))
        for position in unit["positions"]:
            add(position["id"], position["facts"])
    flat: dict[tuple[str, str], dict[str, Any]] = {}
    fact_validator = Draft7Validator(
        {
            "definitions": _load_schema("opportunities")["definitions"],
            "$ref": "#/definitions/Fact",
        }
    )
    for fact in e["facts_flat"]:
        if next(fact_validator.iter_errors(fact), None) is not None:
            _fail("DELIVERY_SCHEMA_INVALID")
        entity = fact.get("entity_id")
        if not isinstance(entity, str) or entity not in entities:
            _fail("FACT_ENTITY_UNKNOWN")
        kind = entities[entity]["kind"]
        if (
            fact.get("opportunity_id", root) != root
            or fact.get("entity_kind", kind) != kind
            or fact.get("level", kind) != kind
        ):
            _fail("FACT_ENTITY_BINDING_INVALID")
        key = _fact_key(entity, fact)
        if key in flat:
            _fail("FACT_FIELD_DUPLICATE")
        flat[key] = fact
    if not nested or not flat:
        _fail("DELIVERY_FACTS_EMPTY")
    if set(nested) != set(flat) or any(
        _fact_payload(fact) != _fact_payload(flat[key]) for key, fact in nested.items()
    ):
        _fail("FACT_REPRESENTATION_MISMATCH")
    return [(entity, fact) for (entity, _), fact in nested.items()]


def _normalized(text: str) -> str:
    return " ".join(text.split())


def _canonical_html_text(text: str) -> str:
    # Whitespace is collapsed, not deleted. Keep Latin word separators, numeric
    # boundaries, punctuation, ZWJ/ZWNJ and bidi controls. No Unicode folding.
    return _normalized(_HAN_FORMAT_MARKS.sub("", text.strip().lstrip("\ufeff")))


def _html_quote_text(root: etree._Element, *, join_cell_wraps: bool = False) -> str:
    """Preserve inline text while separating paragraphs, rows and cells.

    This checks source text structure, not CSS layout or the meaning of a fact.
    The original artifact bytes remain untouched.
    """
    # Only an explicitly selected cell may join adjacent Han text across p/br
    # layout boundaries. None records an extractor-generated separator; literal
    # source whitespace and all row/cell/other block boundaries remain distinct.
    cell_wraps = join_cell_wraps and root.tag in {"td", "th"}
    parts: list[str | None] = []
    for event, node in etree.iterwalk(root, events=("start", "end", "comment", "pi")):
        if node.tag in _HTML_TEXT_BOUNDARIES:
            parts.append(None if cell_wraps and node.tag in {"p", "br"} else " ")
        if event == "start" and isinstance(node.tag, str) and node.text:
            parts.append(node.text)
        elif event != "start" and node is not root and node.tail:
            parts.append(node.tail)
    rendered: list[str] = []
    pending_wrap = False
    for part in parts:
        if part is None:
            pending_wrap = True
        elif part:
            if (
                pending_wrap
                and rendered
                and not (re.fullmatch(_HAN, rendered[-1][-1]) and re.fullmatch(_HAN, part[0]))
            ):
                rendered.append(" ")
            rendered.append(part)
            pending_wrap = False
    return "".join(rendered)


def _quote_support(
    artifact: ValidatedArtifact,
    quote: str,
    locator: dict[str, Any],
    issues: set[str],
    *,
    spreadsheet_cache: dict[str, _SpreadsheetText] | None = None,
) -> bool:
    media = artifact.media_type.partition(";")[0].strip().lower()
    supported_locator = False
    selected: str | None = None
    try:
        if media == "text/html":
            tree = _parse_html_tree(artifact.content, utf8_fallback=True)
            canonical_quote = _canonical_html_text(quote)
            if not canonical_quote:
                _fail("EVIDENCE_QUOTE_INVALID")
            selector = locator.get("selector")
            if isinstance(selector, str) and set(locator) == {"selector"}:
                nodes = tree.cssselect(selector)
                if not nodes:
                    _fail("EVIDENCE_LOCATOR_MISMATCH")
                # A selector may match several locations; never stitch them into a quote.
                texts = [_html_quote_text(node) for node in nodes]
                texts.extend(
                    _html_quote_text(node, join_cell_wraps=True)
                    for node in nodes
                    if node.tag in {"td", "th"}
                )
                supported_locator = True
            else:
                texts = [_html_quote_text(tree)]
                if set(locator) == {"url"}:
                    if locator["url"] != artifact.url:
                        _fail("EVIDENCE_LOCATOR_MISMATCH")
                    supported_locator = True
            if not any(canonical_quote in _canonical_html_text(text) for text in texts):
                _fail("EVIDENCE_QUOTE_MISMATCH")
        elif media == "application/pdf":
            reader = PdfReader(BytesIO(artifact.content), strict=True)
            if reader.is_encrypted:
                issues.add(f"EVIDENCE_FORMAT_REVIEW_REQUIRED:{artifact.artifact_id}")
                return False
            if len(reader.pages) > MAX_PDF_PAGES:
                _fail("EVIDENCE_RESOURCE_LIMIT_EXCEEDED")
            page = locator.get("page")
            if type(page) is int and set(locator) == {"page"}:
                if not 1 <= page <= len(reader.pages):
                    _fail("EVIDENCE_LOCATOR_MISMATCH")
                selected = reader.pages[page - 1].extract_text() or ""
                supported_locator = True
        elif media == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            selected, supported_locator = _spreadsheet_text(
                artifact.content, locator, cache=spreadsheet_cache, cache_key=artifact.artifact_id
            )
        else:
            issues.add(f"EVIDENCE_FORMAT_REVIEW_REQUIRED:{artifact.artifact_id}")
            return False
    except DeliveryValidationError:
        raise
    except ExpectedParseError as error:
        if error.code == "HTML_ENCODING_UNRESOLVED":
            issues.add(f"EVIDENCE_FORMAT_REVIEW_REQUIRED:{artifact.artifact_id}")
            return False
        raise DeliveryValidationError("EVIDENCE_MATERIAL_INVALID") from None
    except Exception:
        # Parser-specific errors must not expose material bytes or local paths.
        raise DeliveryValidationError("EVIDENCE_MATERIAL_INVALID") from None
    if selected is not None and _normalized(quote) not in _normalized(selected):
        _fail("EVIDENCE_QUOTE_MISMATCH")
    if not supported_locator:
        issues.add(f"EVIDENCE_LOCATOR_REVIEW_REQUIRED:{artifact.artifact_id}")
    return supported_locator


def _spreadsheet_text(
    content: bytes,
    locator: dict[str, Any],
    *,
    cache: dict[str, _SpreadsheetText] | None = None,
    cache_key: str = "",
) -> tuple[str | None, bool]:
    sheet = locator.get("sheet")
    row = locator.get("row")
    cell = locator.get("cell")
    column = locator.get("col")
    numeric_cell = set(locator) == {"sheet", "row", "col"}
    if numeric_cell and (
        type(row) is not int
        or not 1 <= row <= 1048576
        or type(column) is not int
        or not 1 <= column <= 16384
    ):
        _fail("EVIDENCE_LOCATOR_MISMATCH")
    if not (
        (set(locator) == {"sheet", "row"} and type(row) is int)
        or (set(locator) == {"sheet", "cell"} and isinstance(cell, str))
        or numeric_cell
    ):
        return None, False
    if "cell" in locator:
        match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", str(cell))
        if match is None:
            _fail("EVIDENCE_LOCATOR_MISMATCH")
        column, row = column_index_from_string(match[1]), int(match[2])
        if column > 16384 or row > 1048576:
            _fail("EVIDENCE_LOCATOR_MISMATCH")
    tables = cache.get(cache_key) if cache is not None else None
    if tables is None:
        tables = _read_spreadsheet_text(content)
        if cache is not None:
            cache[cache_key] = tables
    if not isinstance(sheet, str) or sheet not in tables.rows:
        _fail("EVIDENCE_LOCATOR_MISMATCH")
    if type(row) is not int or not 1 <= row <= tables.row_counts[sheet]:
        _fail("EVIDENCE_LOCATOR_MISMATCH")
    values = tables.rows[sheet].get(row, {})
    if column is not None:
        return values.get(column, ""), True
    return " ".join(values.values()), True


def _read_spreadsheet_text(content: bytes) -> _SpreadsheetText:
    """Read cells once per validation, retaining zeroes and exact worksheet names."""
    try:
        return _bounded_spreadsheet_text(content)
    except ExpectedParseError as error:
        if error.code in {
            "XLSX_ENTRY_LIMIT_EXCEEDED",
            "XLSX_UNCOMPRESSED_SIZE_EXCEEDED",
            "XLSX_EXPANSION_LIMIT_EXCEEDED",
            "XLSX_WORKSHEET_LIMIT_EXCEEDED",
        }:
            _fail("EVIDENCE_RESOURCE_LIMIT_EXCEEDED")
        raise


def _bounded_spreadsheet_text(content: bytes) -> _SpreadsheetText:
    _preflight_archive(content)
    book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.sheetnames) > MAX_WORKSHEETS:
            _fail("EVIDENCE_RESOURCE_LIMIT_EXCEEDED")
        sheets: dict[str, dict[int, dict[int, str]]] = {}
        row_counts: dict[str, int] = {}
        expansion = WorksheetExpansionBudget()
        for worksheet in book.worksheets:
            _preflight_loaded_worksheet(worksheet, expansion)
            _ignore_declared_dimensions(worksheet)
            rows = sheets[worksheet.title] = {}
            row_counts[worksheet.title] = 0
            for number, values in enumerate(worksheet.iter_rows(values_only=True), 1):
                populated = {
                    i: str(value) for i, value in enumerate(values, 1) if value is not None
                }
                if populated:
                    rows[number] = populated
                row_counts[worksheet.title] = number
        return _SpreadsheetText(sheets, row_counts)
    finally:
        book.close()


def validate_delivery(
    result_files: dict[str, bytes], artifact_bytes: dict[str, bytes]
) -> ValidatedDelivery:
    if set(result_files) != _FILES or any(not isinstance(v, bytes) for v in result_files.values()):
        _fail("DELIVERY_FILES_INVALID")
    o = _json(result_files["opportunities.json"])
    e = _json(result_files["evidence.json"])
    _schema("opportunities", o)
    _schema("evidence", e)
    try:
        report = result_files["report.md"].decode("utf-8")
    except UnicodeError:
        raise DeliveryValidationError("DELIVERY_REPORT_INVALID") from None
    if not report.strip():
        _fail("DELIVERY_REPORT_INVALID")
    artifacts = _artifacts(e, artifact_bytes)
    by_id = {artifact.artifact_id: artifact for artifact in artifacts}
    entities, root = _entities(o, e)
    raw_facts = _facts(o, e, entities, root)
    issues = {"HUMAN_FACT_REVIEW_REQUIRED"}
    facts: list[DeliveryFact] = []
    spreadsheet_cache: dict[str, _SpreadsheetText] = {}
    for entity, fact in raw_facts:
        references: list[DeliveryEvidence] = []
        for ref in fact.get("evidence", []):
            artifact = by_id.get(ref["artifact_id"])
            if artifact is None:
                _fail("EVIDENCE_ARTIFACT_UNKNOWN")
            quote = _string(ref["quote"], "EVIDENCE_QUOTE_INVALID")
            declared = ref.get("sha256")
            if declared is not None and (
                not isinstance(declared, str) or declared.lower() != artifact.sha256
            ):
                _fail("EVIDENCE_HASH_MISMATCH")
            locator = ref.get("locator", {})
            verified = _quote_support(
                artifact, quote, locator, issues, spreadsheet_cache=spreadsheet_cache
            )
            references.append(
                DeliveryEvidence(artifact.artifact_id, quote, locator, artifact.sha256, verified)
            )
        facts.append(
            DeliveryFact(
                entity,
                fact["field"],
                fact["value"],
                fact["status"],
                tuple(references),
                note=fact.get("note"),
            )
        )
    manifest = {
        "files": {key: sha256(value).hexdigest() for key, value in result_files.items()},
        "artifacts": {item.artifact_id: item.sha256 for item in artifacts},
    }
    digest = sha256(
        json.dumps(manifest, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    return ValidatedDelivery(o, e, report, artifacts, tuple(facts), tuple(sorted(issues)), digest)


def preflight_manifest(result_files: dict[str, bytes]) -> tuple[tuple[str, str], ...]:
    """Return only task-relative paths; source/task authorization belongs to the caller."""
    if set(result_files) != _FILES or any(not isinstance(v, bytes) for v in result_files.values()):
        _fail("DELIVERY_FILES_INVALID")
    opportunities = _json(result_files["opportunities.json"])
    evidence = _json(result_files["evidence.json"])
    _schema("opportunities", opportunities)
    _schema("evidence", evidence)
    return tuple((identity, local) for _item, identity, _remote, local in _manifest(evidence))
