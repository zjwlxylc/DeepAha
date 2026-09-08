from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    NumberObject,
)

from deepaha.evidence_verification.adapters.defaults import default_registry
from deepaha.evidence_verification.adapters.spreadsheet import SpreadsheetAdapter
from deepaha.evidence_verification.contracts import (
    ArtifactInput,
    Representation,
    VerificationResult,
)
from deepaha.evidence_verification.registry import AdapterRegistry
from deepaha.evidence_verification.verifier import EvidenceVerifier

FIXTURES = Path(__file__).parents[1] / "fixtures" / "documents"


def source(content: bytes, media: str) -> ArtifactInput:
    return ArtifactInput(
        "attachment", media, "https://example.gov/file", sha256(content).hexdigest(), content
    )


def pdf(
    quote: str, locator: dict[str, object], name: str = "minimal-text.pdf"
) -> VerificationResult:
    return EvidenceVerifier(default_registry()).verify(
        source((FIXTURES / name).read_bytes(), "application/pdf"), quote, locator
    )


def test_pdf_quote_stays_in_its_declared_page_and_retains_page_coordinates() -> None:
    result = pdf("page two.", {"page": 2})
    assert result.verdict == "PASS"
    span = result.matches[0].source_spans[0]
    assert (span.origin_id, span.start, span.end) == ("page:2", 14, 23)
    wrong = pdf("page two.", {"page": 1})
    assert (wrong.verdict, wrong.content_support) == ("FAIL", "NOT_FOUND")


def test_missing_pdf_text_layer_cannot_be_declared_quote_error() -> None:
    result = pdf("doctorate", {"page": 1}, "empty-text.pdf")
    assert (result.verdict, result.content_support) == ("UNVERIFIED", "NOT_ESTABLISHED")
    assert "PDF_TEXT_LAYER_MISSING" in result.reason_codes


@pytest.mark.parametrize("locator", [{"page": 0}, {"page": 3}, {"page": True}, {"page": "1"}])
def test_pdf_wrong_or_invalid_page_is_not_rescued(locator: dict[str, object]) -> None:
    assert pdf("page one.", locator).verdict == "FAIL"


def test_pdf_unexecutable_location_reports_content_without_binding() -> None:
    result = pdf("page one.", {"section": "summary", "human_verify": True})
    assert (result.content_support, result.binding, result.verdict) == (
        "FOUND",
        "UNBOUND",
        "UNVERIFIED",
    )


def test_encrypted_pdf_does_not_report_a_quote_error() -> None:
    reader = PdfReader(BytesIO((FIXTURES / "minimal-text.pdf").read_bytes()))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    writer.encrypt("synthetic-password")
    output = BytesIO()
    writer.write(output)
    result = EvidenceVerifier(default_registry()).verify(
        source(output.getvalue(), "application/pdf"), "page one.", {"page": 1}
    )
    assert result.verdict == "UNVERIFIED"
    assert result.reason_codes == ("READER_PDF_ENCRYPTED",)


def workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "岗位 表"
    sheet.append(["doctorate", 0, "=1+1"])
    sheet.append(["master"])
    sheet["D4"] = "other"
    sheet.append(["first", "second"])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def xlsx(quote: str, locator: dict[str, object]) -> VerificationResult:
    return EvidenceVerifier(default_registry()).verify(
        source(workbook_bytes(), SpreadsheetAdapter.media_types[0]), quote, locator
    )


@pytest.mark.parametrize(
    "locator,quote",
    [
        ({"sheet": "岗位 表", "row": 1}, "doctorate 0 =1+1"),
        ({"sheet": "岗位 表", "cell": "B1"}, "0"),
        ({"sheet": "岗位 表", "row": 1, "col": 3}, "=1+1"),
        ({"sheet": "岗位 表", "row": 4, "col": 4}, "other"),
    ],
)
def test_spreadsheet_preserves_zero_formula_and_sparse_coordinates(
    locator: dict[str, object], quote: str
) -> None:
    result = xlsx(quote, locator)
    assert result.content_support == "FOUND"
    assert result.verdict == ("UNVERIFIED" if "=" in quote else "PASS")


@pytest.mark.parametrize(
    "locator,quote",
    [
        ({"sheet": "岗位 表", "row": 2}, "doctorate"),
        ({"sheet": "岗位 表", "cell": "A1"}, "master"),
        ({"sheet": "岗位 表", "cell": "B1"}, "=1+1"),
        ({"sheet": "岗位 表", "row": 5}, "firstsecond"),
        ({"sheet": "岗位表", "row": 1}, "doctorate"),
        ({"sheet": "岗位 表", "cell": "A3"}, "doctorate"),
        ({"sheet": "岗位 表", "cell": "A99"}, "doctorate"),
        ({"sheet": "岗位 表", "row": True}, "doctorate"),
        ({"sheet": "岗位 表", "cell": "XFE1"}, "doctorate"),
    ],
)
def test_spreadsheet_never_rescues_other_rows_columns_or_whitespace(
    locator: dict[str, object], quote: str
) -> None:
    assert xlsx(quote, locator).verdict == "FAIL"


def test_spreadsheet_row_and_cell_views_do_not_duplicate_source_hits() -> None:
    result = xlsx("doctorate", {"description": "first row"})
    assert (result.verdict, result.content_support) == ("UNVERIFIED", "FOUND")
    assert len(result.matches) == 1
    assert result.matches[0].source_spans[0].origin_id == '["岗位 表",1,1]'


def test_spreadsheet_adversarial_actual_coordinate_is_bounded() -> None:
    original, altered = BytesIO(workbook_bytes()), BytesIO()
    with ZipFile(original) as incoming, ZipFile(altered, "w") as outgoing:
        for item in incoming.infolist():
            content = incoming.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'<row r="1">', b'<row r="999999999">')
            outgoing.writestr(item, content)
    result = EvidenceVerifier(default_registry()).verify(
        source(altered.getvalue(), SpreadsheetAdapter.media_types[0]),
        "doctorate",
        {"sheet": "岗位 表", "row": 1},
    )
    assert result.verdict == "UNVERIFIED"
    assert result.reason_codes == ("READER_XLSX_COORDINATE_INVALID",)


def test_actual_spreadsheet_is_read_once_per_run() -> None:
    class CountedAdapter(SpreadsheetAdapter):
        reads = 0

        def read(self, content: bytes) -> Representation:
            self.reads += 1
            return super().read(content)

    adapter = CountedAdapter()
    verifier = EvidenceVerifier(AdapterRegistry((adapter,)))
    artifact = source(workbook_bytes(), adapter.media_types[0])
    for quote, cell in [("doctorate", "A1"), ("0", "B1"), ("=1+1", "C1")]:
        assert verifier.verify(artifact, quote, {"sheet": "岗位 表", "cell": cell}).verdict == (
            "UNVERIFIED" if cell == "C1" else "PASS"
        )
    assert adapter.reads == 1


@pytest.mark.parametrize(
    "kind", ["image", "inline_image", "form_image", "form_inline_image", "vector"]
)
@pytest.mark.parametrize("quote", ["page one.", "unread image text"])
def test_pdf_mixed_visual_and_text_page_is_incomplete_even_with_some_text(
    kind: str, quote: str
) -> None:
    reader = PdfReader(BytesIO((FIXTURES / "minimal-text.pdf").read_bytes()))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    page = writer.pages[0]
    image = DecodedStreamObject()
    image.set_data(b"\x00\x00\x00")
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(1),
            NameObject("/Height"): NumberObject(1),
            NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    command = b"q 10 0 0 10 10 10 cm /Im1 Do Q"
    objects = DictionaryObject({NameObject("/Im1"): writer._add_object(image)})
    if kind in {"inline_image", "form_inline_image"}:
        command = b"q 10 0 0 10 10 10 cm BI /W 1 /H 1 /CS /RGB /BPC 8 ID \x00\x00\x00 EI Q"
        objects = DictionaryObject()
    if kind.startswith("form_"):
        form = DecodedStreamObject()
        form.set_data(command)
        form.update(
            {
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Form"),
                NameObject("/BBox"): ArrayObject(
                    [NumberObject(0), NumberObject(0), NumberObject(100), NumberObject(100)]
                ),
                NameObject("/Resources"): DictionaryObject({NameObject("/XObject"): objects}),
            }
        )
        objects = DictionaryObject({NameObject("/Fm1"): writer._add_object(form)})
        command = b"/Fm1 Do"
    if kind == "vector":
        command = b"10 10 20 20 re f"
        objects = DictionaryObject()
    resources = page["/Resources"]
    assert isinstance(resources, DictionaryObject)
    resources[NameObject("/XObject")] = objects
    old_content = page.get_contents()
    assert old_content is not None
    stream = DecodedStreamObject()
    stream.set_data(old_content.get_data() + b"\n" + command)
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    result = EvidenceVerifier(default_registry()).verify(
        source(output.getvalue(), "application/pdf"), quote, {"page": 1}
    )
    assert result.verdict == "UNVERIFIED"
    assert result.content_support == ("FOUND" if quote == "page one." else "NOT_ESTABLISHED")
    assert "PDF_VISUAL_CONTENT_UNREAD" in result.reason_codes


@pytest.mark.parametrize("kind", ["bad_shared_string", "malformed_content_types"])
def test_malformed_spreadsheet_returns_reading_failure_without_aborting_run(kind: str) -> None:
    original, altered = BytesIO(workbook_bytes()), BytesIO()
    with ZipFile(original) as incoming, ZipFile(altered, "w") as outgoing:
        for item in incoming.infolist():
            content = incoming.read(item.filename)
            if kind == "bad_shared_string" and item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(
                    b't="inlineStr"><is><t>doctorate</t></is>', b't="s"><v>999</v>'
                )
            if kind == "malformed_content_types" and item.filename == "[Content_Types].xml":
                content = content[:-5]
            outgoing.writestr(item, content)
    result = EvidenceVerifier(default_registry()).verify(
        source(altered.getvalue(), SpreadsheetAdapter.media_types[0]),
        "doctorate",
        {"sheet": "岗位 表", "row": 1},
    )
    assert result.verdict == "UNVERIFIED"
    assert result.reason_codes == ("READER_XLSX_PARSE_FAILED",)


@pytest.mark.parametrize("quote", ["5", "=A1+B1"])
def test_formula_and_cached_display_are_not_mistaken_for_complete_literal_cell(quote: str) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append([2, 3, "=A1+B1"])
    output = BytesIO()
    book.save(output)
    changed = BytesIO()
    with ZipFile(BytesIO(output.getvalue())) as original, ZipFile(changed, "w") as target:
        for item in original.infolist():
            content = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<f>A1+B1</f><v></v>", b"<f>A1+B1</f><v>5</v>")
                assert b"<v>5</v>" in content
            target.writestr(item, content)
    result = EvidenceVerifier(default_registry()).verify(
        source(changed.getvalue(), SpreadsheetAdapter.media_types[0]),
        quote,
        {"sheet": "Sheet", "cell": "C1"},
    )
    assert result.verdict == "UNVERIFIED"
    assert result.content_support == ("FOUND" if quote.startswith("=") else "NOT_ESTABLISHED")
    assert "XLSX_DISPLAY_VALUE_UNREAD" in result.reason_codes


def test_formatted_number_scope_is_incomplete_without_poisoning_other_cells() -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append([0.25, "doctorate"])
    sheet["A1"].number_format = "0%"
    output = BytesIO()
    book.save(output)
    verifier = EvidenceVerifier(default_registry())
    artifact = source(output.getvalue(), SpreadsheetAdapter.media_types[0])
    result = verifier.verify(artifact, "25%", {"sheet": "Sheet", "cell": "A1"})
    assert (result.verdict, result.content_support) == ("UNVERIFIED", "NOT_ESTABLISHED")
    assert (
        verifier.verify(artifact, "doctorate", {"sheet": "Sheet", "cell": "B1"}).verdict == "PASS"
    )
    assert (
        verifier.verify(artifact, "doctorate", {"sheet": "Sheet", "row": 1}).verdict == "UNVERIFIED"
    )
