import copy
import importlib
import json
from dataclasses import asdict
from hashlib import sha256
from io import BytesIO
from typing import Any
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter


def _delivery_module() -> Any:
    try:
        return importlib.import_module("deepaha.investigations.delivery")
    except ModuleNotFoundError:
        pytest.fail("the delivery validation boundary has not been implemented")


def _sample() -> tuple[dict[str, Any], dict[str, Any], dict[str, bytes]]:
    content = b'<html><body><p id="terms">Degree: doctorate.</p></body></html>'
    reference = {
        "artifact_id": "notice",
        "quote": "Degree: doctorate.",
        "locator": {"selector": "#terms"},
    }
    fact = {"field": "degree", "value": "doctorate", "status": "CONFIRMED", "evidence": [reference]}
    opportunities = {
        "opportunity_name": "Synthetic notice",
        "publish_unit": "Synthetic issuer",
        "opportunity_type": "recruitment",
        "announcement_level": [],
        "units": [
            {
                "id": "unit",
                "name": "Unit",
                "parent_id": "announcement",
                "unit_level": [],
                "positions": [
                    {
                        "id": "position",
                        "name": "Position",
                        "code": "P1",
                        "facts": [copy.deepcopy(fact)],
                    }
                ],
            }
        ],
    }
    evidence = {
        "artifacts": [
            {
                "artifact_id": "notice",
                "source_url": "https://example.gov/notice",
                "local_path": "artifacts/notice.html",
                "remote_path": "/workspace/task/artifacts/notice.html",
                "file_name": "notice.html",
                "media_type": "text/html",
                "sha256": sha256(content).hexdigest(),
            }
        ],
        "entities": [
            {
                "id": "announcement",
                "name": "Synthetic notice",
                "kind": "announcement",
                "parent_id": None,
            },
            {"id": "unit", "name": "Unit", "kind": "unit", "parent_id": "announcement"},
            {
                "id": "position",
                "name": "Position",
                "kind": "position",
                "code": "P1",
                "parent_id": "unit",
            },
        ],
        "facts_flat": [
            {
                "entity_id": "position",
                "opportunity_id": "announcement",
                "entity_kind": "position",
                "level": "position",
                **copy.deepcopy(fact),
            }
        ],
    }
    return opportunities, evidence, {"notice": content}


def _validate(o: dict[str, Any], e: dict[str, Any], artifacts: dict[str, bytes]) -> Any:
    return _delivery_module().validate_delivery(
        {
            "opportunities.json": json.dumps(o).encode(),
            "evidence.json": json.dumps(e).encode(),
            "report.md": b"# Synthetic report",
        },
        artifacts,
    )


def _fact(o: dict[str, Any]) -> dict[str, Any]:
    return o["units"][0]["positions"][0]["facts"][0]  # type: ignore[no-any-return]


def test_validates_actual_nested_and_flat_shape_against_real_bytes() -> None:
    o, e, artifacts = _sample()
    result = _validate(o, e, artifacts)
    assert result.artifacts[0].content == artifacts["notice"]
    assert result.artifacts[0].sha256 == sha256(artifacts["notice"]).hexdigest()
    assert result.facts[0].entity_id == "position"
    assert result.facts[0].evidence[0].mechanically_verified
    assert "HUMAN_FACT_REVIEW_REQUIRED" in result.issues
    assert result.sha256 == _validate(o, e, artifacts).sha256
    verification = result.facts[0].evidence[0].verification
    assert verification.content_support == "FOUND" and verification.declared_locator == "VERIFIED"
    assert verification.original_locator == {"selector": "#terms"}
    assert verification.reader is not None and verification.representation_sha256


def test_default_detects_ambiguous_locations_without_rewriting_legacy_replay() -> None:
    o, e, artifacts = _sample()
    artifacts["notice"] = (
        b'<main id="terms"><p>Degree: doctorate.</p><p>Degree: doctorate.</p></main>'
    )
    e["artifacts"][0]["sha256"] = sha256(artifacts["notice"]).hexdigest()
    files = {
        "opportunities.json": json.dumps(o).encode(),
        "evidence.json": json.dumps(e).encode(),
        "report.md": b"Synthetic report",
    }
    module = _delivery_module()
    new, old = (
        module.validate_delivery(files, artifacts),
        module.validate_legacy_delivery(files, artifacts),
    )
    assert old.facts[0].evidence[0].mechanically_verified
    assert old.facts[0].evidence[0].verification is None
    assert new.facts[0].evidence[0].verification.binding == "AMBIGUOUS"
    assert not new.facts[0].evidence[0].mechanically_verified
    assert old.sha256 == new.sha256
    assert old.facts[0].evidence[0].quote == new.facts[0].evidence[0].quote


@pytest.mark.parametrize("note", [None, "Candidate inference only; the issuer must clarify."])
def test_fact_notes_survive_validation_for_review(note: str | None) -> None:
    o, e, artifacts = _sample()
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["note"] = note
    result = _validate(o, e, artifacts)
    assert asdict(result.facts[0])["note"] == note
    assert result.facts[0].status == "CONFIRMED"
    assert "HUMAN_FACT_REVIEW_REQUIRED" in result.issues


def test_note_disagreement_between_nested_and_flat_facts_is_rejected() -> None:
    o, e, artifacts = _sample()
    _fact(o)["note"] = "Candidate inference only."
    e["facts_flat"][0]["note"] = "No uncertainty."
    with pytest.raises(
        _delivery_module().DeliveryValidationError, match="FACT_REPRESENTATION_MISMATCH"
    ):
        _validate(o, e, artifacts)


def test_nullable_nested_parent_uses_unambiguous_containment_but_flat_parent_stays_exact() -> None:
    o, e, artifacts = _sample()
    o["units"][0]["parent_id"] = None
    assert _validate(o, e, artifacts).facts[0].entity_id == "position"
    e["entities"][1]["parent_id"] = None
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="ENTITY_PARENT_MISMATCH"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize(
    "mutation,code",
    [
        ("hash", "ARTIFACT_HASH_MISMATCH"),
        ("missing", "ARTIFACT_MISSING"),
        ("duplicate_id", "ARTIFACT_ID_DUPLICATE"),
        ("same_name", "ARTIFACT_PATH_COLLISION"),
        ("entity_duplicate", "ENTITY_ID_DUPLICATE"),
        ("wrong_parent", "ENTITY_PARENT_MISMATCH"),
        ("unknown_entity", "FACT_ENTITY_UNKNOWN"),
        ("flat_value", "FACT_REPRESENTATION_MISMATCH"),
        ("no_evidence", "DELIVERY_SCHEMA_INVALID"),
        ("zero", "DELIVERY_FACTS_EMPTY"),
        ("quote", "EVIDENCE_QUOTE_MISMATCH"),
        ("locator", "EVIDENCE_LOCATOR_MISMATCH"),
        ("ref_hash", "EVIDENCE_HASH_MISMATCH"),
        ("ref_missing", "EVIDENCE_ARTIFACT_UNKNOWN"),
    ],
)
def test_rejects_broken_byte_entity_and_fact_bindings(mutation: str, code: str) -> None:
    o, e, artifacts = _sample()
    if mutation == "hash":
        artifacts["notice"] += b"changed"
    elif mutation == "missing":
        artifacts.clear()
    elif mutation in {"duplicate_id", "same_name"}:
        added = copy.deepcopy(e["artifacts"][0])
        if mutation == "same_name":
            added["artifact_id"] = "another"
            artifacts["another"] = artifacts["notice"]
        e["artifacts"].append(added)
    elif mutation == "entity_duplicate":
        e["entities"].append(copy.deepcopy(e["entities"][0]))
    elif mutation == "wrong_parent":
        e["entities"][2]["parent_id"] = "announcement"
    elif mutation == "unknown_entity":
        e["facts_flat"][0]["entity_id"] = "ghost"
    elif mutation == "flat_value":
        e["facts_flat"][0]["value"] = "bachelor"
    elif mutation == "zero":
        o["units"][0]["positions"][0]["facts"] = []
        e["facts_flat"] = []
    else:
        for fact in (_fact(o), e["facts_flat"][0]):
            if mutation == "no_evidence":
                fact["evidence"] = []
            elif mutation == "quote":
                fact["evidence"][0]["quote"] = "Salary: one million."
            elif mutation == "locator":
                fact["evidence"][0]["locator"] = {"selector": "#wrong"}
            elif mutation == "ref_hash":
                fact["evidence"][0]["sha256"] = "0" * 64
            elif mutation == "ref_missing":
                fact["evidence"][0]["artifact_id"] = "ghost"
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError) as error:
        _validate(o, e, artifacts)
    assert error.value.code == code


@pytest.mark.parametrize(
    "path",
    [
        "../notice.html",
        "/tmp/notice.html",
        "C:/notice.html",
        "artifacts/../notice.html",
        "artifacts\\notice.html",
        "artifacts/notice.html:stream",
        "artifacts/CON",
    ],
)
def test_never_accepts_local_path_escape_or_platform_alias(path: str) -> None:
    o, e, artifacts = _sample()
    e["artifacts"][0]["local_path"] = path
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="ARTIFACT_PATH_INVALID"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize(
    "payload", [b'{"units":[],"units":[]}', b'{"value":NaN}', b'{"value":Infinity}']
)
def test_rejects_ambiguous_or_nonfinite_json(payload: bytes) -> None:
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="DELIVERY_JSON_INVALID"):
        module.validate_delivery(
            {"opportunities.json": payload, "evidence.json": b"{}", "report.md": b"report"}, {}
        )


def test_preserves_non_utf8_original_and_marks_unsupported_format_for_review() -> None:
    o, e, artifacts = _sample()
    artifacts["notice"] = b"\xd0\xcf\x11\xe0\xff\x81binary"
    artifact = e["artifacts"][0]
    artifact.update(
        media_type="application/msword",
        local_path="artifacts/notice.doc",
        remote_path="/workspace/task/artifacts/notice.doc",
        file_name="notice.doc",
        sha256=sha256(artifacts["notice"]).hexdigest(),
    )
    result = _validate(o, e, artifacts)
    assert result.artifacts[0].content == b"\xd0\xcf\x11\xe0\xff\x81binary"
    assert not result.facts[0].evidence[0].mechanically_verified
    assert "EVIDENCE_FORMAT_REVIEW_REQUIRED:notice" in result.issues


def test_section_locator_is_not_misrepresented_as_mechanically_verified() -> None:
    o, e, artifacts = _sample()
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["locator"] = {"section": "Requirements"}
    result = _validate(o, e, artifacts)
    assert not result.facts[0].evidence[0].mechanically_verified
    assert "EVIDENCE_LOCATOR_REVIEW_REQUIRED:notice" in result.issues


@pytest.mark.parametrize(
    "markup,selector,quote,accepted",
    [
        (
            '<p id="terms">公共管理（<span>125200</span>）</p>',
            "#terms",
            "公共管理（125200）",
            True,
        ),
        (
            '<p id="terms"><span>博士</span><strong>研究生</strong></p>',
            "#terms",
            "博士研究生",
            True,
        ),
        ('<p id="terms">Degree: <b>doctorate</b>.</p>', "#terms", "Degree: doctorate.", True),
        ('<p id="terms">Degree: <b>doctorate</b>.</p>', "#terms", "Degree:doctorate.", False),
        ('<p id="terms">博<script>ignore</script>士</p>', "#terms", "博士", True),
        ('<p id="terms">博<!-- hidden -->士</p>', "#terms", "博士", True),
        (
            '<table><tr><td id="terms"><p>学历、学</p><p>位要求</p></td></tr></table>',
            "#terms",
            "学历、学位要求",
            True,
        ),
        (
            '<table><tr><th id="terms"><p><span>学历、学</span></p>'
            "<p><strong>位要求</strong></p></th></tr></table>",
            "#terms",
            "学历、学 位要求",
            True,
        ),
        (
            '<table><tr><td id="terms"><p>学历、学 </p><p>位要求</p></td></tr></table>',
            "#terms",
            "学历、学位要求",
            False,
        ),
        (
            '<table><tr><td id="terms"><p>1</p><p>2</p></td></tr></table>',
            "#terms",
            "12",
            False,
        ),
        (
            '<table><tr><td id="terms"><p>not</p><p>eligible</p></td></tr></table>',
            "#terms",
            "noteligible",
            False,
        ),
        (
            '<table id="terms"><tr><td>学历、学</td><td>位要求</td></tr></table>',
            "#terms",
            "学历、学位要求",
            False,
        ),
        (
            '<table><tr><td id="terms"><p>须</p><p>持证</p><p>方可报名</p></td></tr></table>',
            "#terms",
            "须方可报名",
            False,
        ),
        (
            '<table><tr><td id="terms"><p>学历：本科</p><p>须持证</p></td></tr></table>',
            "#terms",
            "学历：本科无需持证",
            False,
        ),
        ('<p id="terms">博士\u200b研究生</p>', "#terms", "博士研究生", True),
        ('<p id="terms">博士研究生</p>', "#terms", "\ufeff博士研究生", True),
        ('<p id="terms">Degree:\r\n\t&nbsp;doctorate.</p>', "#terms", "Degree: doctorate.", True),
        ('<p id="terms">not eligible</p>', "#terms", "noteligible", False),
        ('<p id="terms">a\u200cb</p>', "#terms", "ab", False),
        ('<p id="terms">2026-12-31</p>', "#terms", "2026-12-30", False),
        ('<p id="terms">博士</p>其他要求', "#terms", "博士其他要求", False),
        (
            '<div id="terms"><p>博士</p><p>研究生</p></div>',
            "#terms",
            "博士研究生",
            False,
        ),
        (
            '<div id="terms"><p>博士</p><p>研究生</p></div>',
            "#terms",
            "博士 研究生",
            True,
        ),
        (
            '<div id="terms"><p>第一条</p><p>中间条款</p><p>第三条</p></div>',
            "#terms",
            "第一条……第三条",
            False,
        ),
        (
            '<p id="first">第一条</p><p>中间条款</p><p id="last">第三条</p>',
            "#first, #last",
            "第一条 第三条",
            False,
        ),
        (
            '<p id="first">第一条</p><p>中间条款</p><p id="last">第三条</p>',
            "#first, #last",
            "第三条",
            True,
        ),
        (
            '<table id="terms"><tr><td>岗位</td><td>人数</td></tr>'
            "<tr><td>甲</td><td>1</td></tr><tr><td>乙</td><td>2</td></tr></table>",
            "#terms",
            "岗位 人数 乙 2",
            False,
        ),
    ],
)
def test_html_quotes_preserve_inline_text_without_accepting_omitted_passages(
    markup: str, selector: str, quote: str, accepted: bool
) -> None:
    o, e, artifacts = _sample()
    artifacts["notice"] = (
        '<html><head><meta charset="utf-8"></head><body>' + markup + "</body></html>"
    ).encode()
    e["artifacts"][0]["sha256"] = sha256(artifacts["notice"]).hexdigest()
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0].update(quote=quote, locator={"selector": selector})
    if accepted:
        result = _validate(o, e, artifacts)
        assert result.facts[0].evidence[0].mechanically_verified
        assert result.facts[0].evidence[0].quote == quote
        assert result.artifacts[0].content == artifacts["notice"]
        assert "HUMAN_FACT_REVIEW_REQUIRED" in result.issues
    else:
        with pytest.raises(
            _delivery_module().DeliveryValidationError, match="EVIDENCE_QUOTE_MISMATCH"
        ):
            _validate(o, e, artifacts)


def test_xlsx_quote_is_checked_in_the_declared_row_not_another_position() -> None:
    o, e, artifacts = _sample()
    book = Workbook()
    book.active.append(["Position", "Degree"])  # type: ignore[union-attr]
    book.active.append(["P1", "doctorate"])  # type: ignore[union-attr]
    book.active.append(["P2", "bachelor"])  # type: ignore[union-attr]
    output = BytesIO()
    book.save(output)
    artifacts["notice"] = output.getvalue()
    e["artifacts"][0].update(
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        local_path="artifacts/notice.xlsx",
        remote_path="/workspace/task/artifacts/notice.xlsx",
        file_name="notice.xlsx",
        sha256=sha256(artifacts["notice"]).hexdigest(),
    )
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0].update(quote="doctorate", locator={"sheet": "Sheet", "row": 2})
    assert _validate(o, e, artifacts).facts[0].evidence[0].mechanically_verified
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["locator"]["row"] = 3
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="EVIDENCE_QUOTE_MISMATCH"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize(
    "col,quote,valid", [(2, "doctorate", True), (1, "doctorate", False), (3, "0", True)]
)
def test_xlsx_numeric_column_locator_checks_exact_cell(col: int, quote: str, valid: bool) -> None:
    module = _delivery_module()
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["P1", "doctorate", 0])
    output = BytesIO()
    book.save(output)
    book.close()
    artifact = module.ValidatedArtifact(
        "table",
        "https://example.gov/table.xlsx",
        "artifacts/table.xlsx",
        sha256(output.getvalue()).hexdigest(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        output.getvalue(),
    )
    if valid:
        assert module._quote_support(
            artifact, quote, {"sheet": "Sheet", "row": 1, "col": col}, set()
        )
    else:
        with pytest.raises(module.DeliveryValidationError, match="EVIDENCE_QUOTE_MISMATCH"):
            module._quote_support(artifact, quote, {"sheet": "Sheet", "row": 1, "col": col}, set())


def test_utf8_fragment_is_validated_with_same_text_as_document_parser() -> None:
    o, e, artifacts = _sample()
    artifacts["notice"] = '<p id="terms">报考要求：<strong>本科</strong>及以上。</p>'.encode()
    e["artifacts"][0]["sha256"] = sha256(artifacts["notice"]).hexdigest()
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["quote"] = "本科及以上"
    assert _validate(o, e, artifacts).facts[0].evidence[0].mechanically_verified


def test_many_xlsx_references_read_each_workbook_once_and_do_not_reuse_other_deliveries(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from deepaha.documents import spreadsheet

    original = load_workbook
    reads = 0

    def counted(*args: Any, **kwargs: Any) -> Any:
        nonlocal reads
        reads += 1
        return original(*args, **kwargs)

    monkeypatch.setattr(spreadsheet, "load_workbook", counted)
    o, e, artifacts = _sample()
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["doctorate", "P1"])
    output = BytesIO()
    book.save(output)
    book.close()
    artifacts["notice"] = output.getvalue()
    e["artifacts"][0].update(
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        sha256=sha256(artifacts["notice"]).hexdigest(),
    )
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"] = [
            {
                "artifact_id": "notice",
                "quote": "doctorate",
                "locator": {"sheet": "Sheet", "cell": "A1"},
            },
            {
                "artifact_id": "notice",
                "quote": "P1",
                "locator": {"sheet": "Sheet", "row": 1, "col": 2},
            },
        ]
    first = _validate(o, e, artifacts)
    assert all(r.mechanically_verified for r in first.facts[0].evidence)
    assert reads == 1
    assert _validate(o, e, artifacts).sha256 == first.sha256
    assert reads == 2


def test_preflight_returns_only_validated_task_relative_artifact_paths() -> None:
    o, e, _ = _sample()
    module = _delivery_module()
    preflight = getattr(module, "preflight_manifest", None)
    assert callable(preflight), "download must have a path validation boundary"
    files = {
        "opportunities.json": json.dumps(o).encode(),
        "evidence.json": json.dumps(e).encode(),
        "report.md": b"report",
    }
    assert preflight(files) == (("notice", "artifacts/notice.html"),)
    e["artifacts"][0]["local_path"] = "/workspace/another-task/artifacts/notice.html"
    files["evidence.json"] = json.dumps(e).encode()
    with pytest.raises(module.DeliveryValidationError, match="ARTIFACT_PATH_INVALID"):
        preflight(files)


def test_preflight_requires_nonempty_manifest_before_any_download() -> None:
    o, e, _ = _sample()
    e["artifacts"] = []
    module = _delivery_module()
    preflight = getattr(module, "preflight_manifest", None)
    assert callable(preflight), "download must have a path validation boundary"
    with pytest.raises(module.DeliveryValidationError, match="ARTIFACT_MANIFEST_EMPTY"):
        preflight(
            {
                "opportunities.json": json.dumps(o).encode(),
                "evidence.json": json.dumps(e).encode(),
                "report.md": b"report",
            }
        )


def test_json_exponent_overflow_cannot_bypass_nonfinite_rejection() -> None:
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="DELIVERY_JSON_INVALID"):
        module.validate_delivery(
            {
                "opportunities.json": b'{"number":1e999}',
                "evidence.json": b"{}",
                "report.md": b"report",
            },
            {},
        )


def test_malformed_flat_evidence_is_rejected_as_schema_error() -> None:
    o, e, artifacts = _sample()
    e["facts_flat"][0]["evidence"] = "not an evidence array"
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="DELIVERY_SCHEMA_INVALID"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize("path", ["artifacts/%252e%252e/notice.html", "artifacts/notice?.html"])
def test_path_cannot_change_meaning_after_url_or_windows_processing(path: str) -> None:
    o, e, artifacts = _sample()
    e["artifacts"][0]["local_path"] = path
    e["artifacts"][0]["remote_path"] = path
    e["artifacts"][0]["file_name"] = path.rsplit("/", 1)[-1]
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="ARTIFACT_PATH_INVALID"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize(
    "media,locator",
    [
        ("application/pdf", {"page": 1}),
        (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            {"sheet": "S", "row": 1},
        ),
    ],
)
def test_corrupted_supported_material_has_stable_validation_error(
    media: str,
    locator: dict[str, Any],
) -> None:
    o, e, artifacts = _sample()
    artifacts["notice"] = b"corrupt original"
    e["artifacts"][0].update(media_type=media, sha256=sha256(artifacts["notice"]).hexdigest())
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["locator"] = locator
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="EVIDENCE_MATERIAL_INVALID"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize("kind", ["pdf_pages", "xlsx_sheets", "xlsx_declared_expansion"])
def test_material_resource_limits_fail_before_unbounded_extraction(kind: str) -> None:
    o, e, artifacts = _sample()
    output = BytesIO()
    locator: dict[str, Any]
    if kind == "pdf_pages":
        writer = PdfWriter()
        for _ in range(501):
            writer.add_blank_page(width=10, height=10)
        writer.write(output)
        media, locator = "application/pdf", {"page": 1}
    else:
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        locator = {"sheet": "Sheet", "row": 1}
        if kind == "xlsx_sheets":
            book = Workbook()
            for number in range(128):
                book.create_sheet(f"Sheet{number}")
            book.save(output)
        else:
            # A tiny archive can declare dangerous expansion; do not allocate that expansion.
            with ZipFile(output, "w") as archive:
                archive.writestr("xl/worksheets/sheet1.xml", b"small")
            raw = bytearray(output.getvalue())
            central = raw.index(b"PK\x01\x02")
            raw[central + 24 : central + 28] = (100_000_001).to_bytes(4, "little")
            output = BytesIO(raw)
    artifacts["notice"] = output.getvalue()
    e["artifacts"][0].update(media_type=media, sha256=sha256(artifacts["notice"]).hexdigest())
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["locator"] = locator
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="EVIDENCE_RESOURCE_LIMIT_EXCEEDED"):
        _validate(o, e, artifacts)


@pytest.mark.parametrize(
    "identity", ["a/b", ".", "..", "a\\b", "a b", "a%2fb", "原件", "", "a" * 257]
)
@pytest.mark.parametrize("preflight", [True, False])
def test_artifact_identity_must_be_a_safe_download_path_segment(
    identity: str,
    preflight: bool,
) -> None:
    o, e, artifacts = _sample()
    e["artifacts"][0]["artifact_id"] = identity
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["artifact_id"] = identity
    artifacts = {identity: artifacts["notice"]}
    module = _delivery_module()
    with pytest.raises(module.DeliveryValidationError, match="ARTIFACT_ID_INVALID"):
        if preflight:
            module.preflight_manifest(
                {
                    "opportunities.json": json.dumps(o).encode(),
                    "evidence.json": json.dumps(e).encode(),
                    "report.md": b"report",
                }
            )
        else:
            _validate(o, e, artifacts)


@pytest.mark.parametrize("identity", ["a", "art_notice-v1.2", "a" * 256])
def test_safe_artifact_identity_and_maximum_length_remain_downloadable(identity: str) -> None:
    o, e, artifacts = _sample()
    e["artifacts"][0]["artifact_id"] = identity
    for fact in (_fact(o), e["facts_flat"][0]):
        fact["evidence"][0]["artifact_id"] = identity
    result = _validate(o, e, {identity: artifacts["notice"]})
    assert result.artifacts[0].artifact_id == identity
