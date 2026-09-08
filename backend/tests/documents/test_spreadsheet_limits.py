from hashlib import sha256
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from deepaha.documents.parser import ExpectedParseError
from deepaha.documents.spreadsheet import OpenpyxlSpreadsheetParser, _preflight_archive
from deepaha.investigations.delivery import (
    DeliveryValidationError,
    ValidatedArtifact,
    _quote_support,
)


def _workbook(*sheets: str) -> bytes:
    book = Workbook()
    for _ in sheets[1:]:
        book.create_sheet()
    original = BytesIO()
    book.save(original)
    book.close()
    output = BytesIO()
    with ZipFile(original) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            content = source.read(entry)
            for index, rows in enumerate(sheets, 1):
                if entry.filename == f"xl/worksheets/sheet{index}.xml":
                    content = (
                        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                        f'<dimension ref="A1:XFD1048576"/><sheetData>{rows}</sheetData></worksheet>'
                    ).encode()
            target.writestr(entry, content)
    return output.getvalue()


@pytest.mark.parametrize(
    "rows",
    [
        '<row r="1000000000"><c r="A1000000000"><v>1</v></c></row>',
        '<row r="1048577"><c r="A1048577"><v>1</v></c></row>',
        '<row r="1"><c r="XFE1"><v>1</v></c></row>',
        '<row r="0"><c r="A0"><v>1</v></c></row>',
        '<row r="-1"><c r="A1"><v>1</v></c></row>',
        '<row r="1"><c r="A2"><v>1</v></c></row>',
        '<row r="2"/><row r="1"/>',
        '<row r="1"><c r="B1"/><c r="A1"/></row>',
    ],
)
def test_invalid_actual_coordinates_rejected_before_grid_expansion(
    rows: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    def never_iterate(*args: object, **kwargs: object) -> None:
        pytest.fail("untrusted worksheet reached grid expansion")

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", never_iterate)
    content = _workbook(rows)
    with pytest.raises(ExpectedParseError, match="XLSX_COORDINATE_INVALID"):
        OpenpyxlSpreadsheetParser().parse(content, artifact_sha256=sha256(content).hexdigest())


def test_expansion_budget_is_cumulative_across_worksheets() -> None:
    rows = '<row r="1048576"><c r="A1048576"><v>0</v></c></row>'
    # Tiny XML, but two sheets would generate over two million row iterations.
    with pytest.raises(ExpectedParseError, match="XLSX_EXPANSION_LIMIT_EXCEEDED"):
        _preflight_archive(_workbook(rows, rows))


def test_column_padding_counts_towards_expansion_budget() -> None:
    rows = "".join(f'<row r="{r}"><c r="XFD{r}"><v>0</v></c></row>' for r in range(1, 124))
    with pytest.raises(ExpectedParseError, match="XLSX_EXPANSION_LIMIT_EXCEEDED"):
        _preflight_archive(_workbook(rows))


def test_sparse_valid_coordinates_zeroes_and_formulas_are_preserved() -> None:
    content = _workbook(
        '<row r="3"><c r="C3"><v>0</v></c></row>'
        '<row><c r="A4"><f>C3+1</f><v>1</v></c><c><v>2</v></c></row>'
    )
    parsed = OpenpyxlSpreadsheetParser().parse(content, artifact_sha256=sha256(content).hexdigest())
    assert parsed.normalized_text == "[Sheet]\nC3\t0\nA4\t=C3+1\nB4\t2\n"


def test_delivery_gate_rejects_expansion_without_marking_it_verified() -> None:
    rows = '<row r="1048576"><c r="A1048576"><v>0</v></c></row>'
    content = _workbook(rows, rows)
    artifact = ValidatedArtifact(
        "table",
        "https://example.gov/table.xlsx",
        "artifacts/table.xlsx",
        sha256(content).hexdigest(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        content,
    )
    with pytest.raises(DeliveryValidationError, match="EVIDENCE_RESOURCE_LIMIT_EXCEEDED"):
        _quote_support(artifact, "0", {"sheet": "Sheet", "row": 1048576}, set())
