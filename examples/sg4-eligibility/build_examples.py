from pathlib import Path
import io,json,zipfile
from openpyxl import Workbook
ROOT=Path(__file__).parent
URL='https://research.example.org/eligibility/'

def make_zip(name,root,files):
 out=ROOT/name
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  z.writestr('opportunities.json',json.dumps(root,ensure_ascii=False,indent=2).encode())
  z.writestr('evidence.json',json.dumps({'artifacts':files.pop('_artifacts')},ensure_ascii=False,indent=2).encode())
  z.writestr('report.md',('# '+root['opportunity_name']+'\n本包为SG4本地虚构验收样本。').encode())
  for k,v in files.items():z.writestr(k,v)

def text_fact(field,value,quote):return {'field':field,'value':value,'status':'CONFIRMED','evidence':[{'artifact_id':'notice','quote':quote,'locator':{'line':1}}]}
def base(title,code,facts,coverage=None,announcement=None):
 r={'opportunity_name':title,'publish_unit':'海湾研究中心（虚构）','opportunity_type':'RESEARCH_PROGRAM','official_url':URL+code,'announcement_level':announcement or [],'units':[{'id':'g1','name':'青年研究计划','positions':[{'id':'p1','name':title+' · 具体项目','code':code,'facts':facts}]}]}
 if coverage:r['metadata']={'qualification_coverage':coverage}
 return r

def text_pack(filename,root,lines):
 body='\n'.join(lines)+'\n';make_zip(filename,root,{'_artifacts':[{'artifact_id':'notice','local_path':'artifacts/notice.txt','url':root['official_url']}],'artifacts/notice.txt':body.encode()})

text_pack('01_eligible_complete.zip',base('本科生研究实践计划','E01',[
 text_fact('学历','本科及以上','学历：本科及以上'),text_fact('毕业届别','2027届毕业生','毕业届别：2027届毕业生'),text_fact('户籍要求','浙江省户籍','户籍要求：浙江省户籍')],coverage='COMPLETE'),['学历：本科及以上','毕业届别：2027届毕业生','户籍要求：浙江省户籍'])

wb=Workbook();ws=wb.active;ws.title='岗位信息表';ws['K5']='硕士研究生以上';ws['M5']='理学（07）、工学（08）';buf=io.BytesIO();wb.save(buf)
r=base('高阶研究助理计划','I01',[
 {'field':'学历','value':'硕士研究生以上','status':'CONFIRMED','evidence':[{'artifact_id':'plan','quote':'硕士研究生以上','locator':{'sheet':'岗位信息表','row':5,'column':'K'}}]},
 {'field':'专业','value':'理学（07）、工学（08）','status':'CONFIRMED','evidence':[{'artifact_id':'plan','quote':'理学（07）、工学（08）','locator':{'sheet':'岗位信息表','row':5,'column':'M'}}]},
])
make_zip('02_ineligible_xlsx.zip',r,{'_artifacts':[{'artifact_id':'plan','local_path':'artifacts/plan.xlsx','url':r['official_url']}],'artifacts/plan.xlsx':buf.getvalue()})

text_pack('03_uncertain_exception.zip',base('青年专项研究计划','U01',[
 text_fact('学历','硕士研究生以上；具有高级职称者可放宽','学历：硕士研究生以上；具有高级职称者可放宽'),
 text_fact('其他条件','须满足项目组另行公布的专项条件','其他条件：须满足项目组另行公布的专项条件')],coverage='COMPLETE'),['学历：硕士研究生以上；具有高级职称者可放宽','其他条件：须满足项目组另行公布的专项条件'])

text_pack('04_likely_bounded.zip',base('传播创新实践计划','L01',[
 text_fact('学历','本科及以上','学历：本科及以上'),text_fact('毕业届别','2027届毕业生','毕业届别：2027届毕业生')]),['学历：本科及以上','毕业届别：2027届毕业生'])
print('built',len(list(ROOT.glob('*.zip'))),'examples')
